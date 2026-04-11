from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook


def write_announcements_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "announcements"
    ws["A1"] = "Текст"
    ws["A2"] = "// Каждая строка — строка объявления (внутри блока)."
    ws["A3"] = "// Блоки объявлений разделяются строкой: !!!"
    ws["A4"] = "// Строки-комментарии начинаются с // или # и игнорируются."
    ws["A5"] = ""
    ws["A6"] = "Сегодня педсовет в 15:00"
    ws["A7"] = "Линейка в понедельник"
    ws["A8"] = "Проверить сменную обувь"
    ws["A9"] = "!!!"
    ws["A10"] = "12 апреля — День космонавтики"
    ws["A11"] = "Классный час по расписанию"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_marquee_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "marquee"
    ws["A1"] = "Текст"
    ws["A2"] = "// 1 строка = 1 сообщение бегущей строки."
    ws["A3"] = "// Комментарии начинаются с // или # и игнорируются."
    ws["A4"] = ""
    ws["A5"] = "Внимание! Идет настройка экрана"
    ws["A6"] = "Сегодня педсовет в 15:00"
    ws["A7"] = "12 апреля — День космонавтики"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    import_dir = root / "data" / "import"
    # шаблоны
    write_announcements_template(import_dir / "announcements_template.xlsx")
    write_marquee_template(import_dir / "marquee_template.xlsx")
    # рабочие файлы для авто-импорта (чтобы «заполнил шаблон — сразу заработало»)
    write_announcements_template(import_dir / "announcements.xlsx")
    write_marquee_template(import_dir / "marquee.xlsx")


if __name__ == "__main__":
    main()

