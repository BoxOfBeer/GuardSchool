"""Минимальные .xlsx-образцы для ручного импорта в админке (кнопки «Скачать образец»)."""
from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
import tempfile

from fastapi import HTTPException
from openpyxl import Workbook

from .gs_admin_http import admin_msg
from .gs_weekly_template import write_weekly_schedule_template_excel


def _weekly_template_bytes() -> bytes:
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        path = Path(tmp.name)
    try:
        write_weekly_schedule_template_excel(path)
        return path.read_bytes()
    finally:
        path.unlink(missing_ok=True)


def _dated_schedule_sample_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Расписание"
    headers = ["Дата", "Класс"] + [f"Урок{i}" for i in range(1, 5)]
    ws.append(headers)
    ws.append([date(2025, 9, 1), "5А", "Математика", "Русский язык", "Литература", ""])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _holidays_sample_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Праздники"
    ws.append(["Дата", "Название", "Описание"])
    ws.append([date(2025, 5, 9), "День Победы", "Выходной день"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _announcements_sample_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "announcements"
    ws["A1"] = "Текст"
    ws["A2"] = "Сегодня педсовет в 15:00"
    ws["A3"] = "Проверить сменную обувь"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _marquee_sample_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "marquee"
    ws["A1"] = "Текст"
    ws["A2"] = "Внимание! Идёт настройка экрана"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_excel_sample_bytes(*, kind: str, lang: str) -> tuple[bytes, str]:
    k = (kind or "").strip().lower()
    if k == "dated":
        return _dated_schedule_sample_bytes(), "schedule_dated_sample.xlsx"
    if k in ("full", "weekly"):
        return _weekly_template_bytes(), "full_schedule_sample.xlsx"
    if k == "sample":
        return _weekly_template_bytes(), "schedule_sample_example.xlsx"
    if k == "holidays":
        return _holidays_sample_bytes(), "holidays_sample.xlsx"
    if k == "announcements":
        return _announcements_sample_bytes(), "announcements_sample.xlsx"
    if k == "marquee":
        return _marquee_sample_bytes(), "marquee_sample.xlsx"
    raise HTTPException(
        status_code=400,
        detail=admin_msg(
            lang,
            "Неизвестный тип образца.",
            "Unknown sample type.",
        ),
    )
