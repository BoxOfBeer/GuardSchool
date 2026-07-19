"""Парсинг Excel (расписание, праздники, объявления, бегущая строка) и автоимпорт из data/import/."""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook

from .gs_admin_http import admin_msg
from .gs_class_key import class_name_from_excel, normalize_class
from .gs_import_state import import_signature, load_import_state, save_import_state
from .gs_jsonio import write_json
from .gs_paths import (
    ANNOUNCEMENTS_PATH,
    AUTO_ANNOUNCEMENTS_IMPORT_PATH,
    AUTO_FULL_SCHEDULE_IMPORT_PATH,
    AUTO_HOLIDAYS_IMPORT_PATH,
    AUTO_MARQUEE_IMPORT_PATH,
    AUTO_SCHEDULE_IMPORT_PATH,
    AUTO_SCHEDULE_SAMPLE_IMPORT_PATH,
    FULL_SCHEDULE_PATH,
    HOLIDAYS_PATH,
    MARQUEE_PATH,
    SCHEDULE_PATH,
    SCHEDULE_SAMPLE_PATH,
)

_WEEKDAY_ALIASES: dict[str, int] = {
    "понедельник": 0,
    "пн": 0,
    "mon": 0,
    "monday": 0,
    "вторник": 1,
    "вт": 1,
    "tue": 1,
    "tues": 1,
    "tuesday": 1,
    "среда": 2,
    "ср": 2,
    "wed": 2,
    "wednesday": 2,
    "четверг": 3,
    "чт": 3,
    "thu": 3,
    "thursday": 3,
    "пятница": 4,
    "пт": 4,
    "fri": 4,
    "friday": 4,
    "суббота": 5,
    "сб": 5,
    "sat": 5,
    "saturday": 5,
    "воскресенье": 6,
    "вс": 6,
    "sun": 6,
    "sunday": 6,
}


def parse_weekday_value(raw: Any) -> int | None:
    """Понедельник=0 … воскресенье=6, как :meth:`datetime.date.weekday`."""
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return int(raw.date().weekday())
    if isinstance(raw, date):
        return int(raw.weekday())
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        n = int(raw)
        if 0 <= n <= 6:
            return n
        if 1 <= n <= 7:
            return (n - 1) % 7
        return None
    s = str(raw).strip()
    if not s:
        return None
    if re.fullmatch(r"\d{1,2}", s):
        n = int(s)
        if 0 <= n <= 6:
            return n
        if 1 <= n <= 7:
            return (n - 1) % 7
        return None
    return _WEEKDAY_ALIASES.get(s.lower().replace("ё", "е"))


def parse_lesson_column_index(header: str) -> int | None:
    """Номер слота из заголовка: «Урок1», «Слот2», «Slot3», «Lesson4» (без регистра)."""
    s = str(header or "").strip()
    if not s:
        return None
    sl = s.lower().replace("ё", "е")
    for prefix in ("урок", "слот", "slot", "lesson"):
        if not sl.startswith(prefix):
            continue
        tail = sl[len(prefix) :].strip()
        if not tail:
            return None
        m = re.match(r"^(\d+)", tail)
        if not m:
            return None
        try:
            n = int(m.group(1))
        except ValueError:
            return None
        return n if 1 <= n <= 24 else None
    return None


def _normalize_excel_header(header: str) -> str:
    return str(header or "").strip().lower().replace("ё", "е")


def _column_index_by_header_aliases(headers: list[str], aliases: frozenset[str]) -> int | None:
    for i, h in enumerate(headers):
        if _normalize_excel_header(h) in aliases:
            return i
    return None


_CLASS_COLUMN_ALIASES = frozenset({"класс", "группа", "class", "group"})
_DATE_COLUMN_ALIASES = frozenset({"дата", "date"})


def _is_blank_excel_scalar(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


def _resolve_class_cell_value(sheet: Any, excel_row: int, class_col_idx: int, values_only_val: Any) -> Any:
    """Для объединённых ячеек values_only даёт None — подставляем значение с якоря диапазона."""
    if not _is_blank_excel_scalar(values_only_val):
        return values_only_val
    col = class_col_idx + 1
    cell = sheet.cell(row=excel_row, column=col)
    if type(cell).__name__ != "MergedCell":
        return cell.value
    for rng in sheet.merged_cells.ranges:
        if cell.coordinate in rng:
            return sheet.cell(row=rng.min_row, column=rng.min_col).value
    return values_only_val


def _pad_row_to_width(row: tuple[Any, ...] | list[Any] | None, width: int) -> list[Any]:
    if row is None:
        return [None] * width
    lst = list(row)
    if len(lst) < width:
        lst.extend([None] * (width - len(lst)))
    return lst[:width]


def _column_has_data_in_rows(data_rows: list[list[Any]], col_idx: int) -> bool:
    for row in data_rows:
        if col_idx >= len(row):
            continue
        v = row[col_idx]
        if v is None:
            continue
        if str(v).strip() != "":
            return True
    return False


def _discover_lesson_specs(
    headers: list[str],
    class_col_idx: int,
    data_rows: list[list[Any]],
) -> list[tuple[int, int]]:
    n = len(headers)
    explicit: list[tuple[int, int]] = []
    for j in range(class_col_idx + 1, n):
        lix = parse_lesson_column_index(headers[j])
        if lix is not None:
            explicit.append((j, lix))
    explicit.sort(key=lambda x: x[0])
    used_j = {j for j, _ in explicit}
    last_j = max((j for j, _ in explicit), default=class_col_idx)
    next_num = max((lix for _, lix in explicit), default=0) + 1

    inferred: list[tuple[int, int]] = []
    for j in range(last_j + 1, n):
        if j in used_j:
            continue
        if headers[j].strip() != "":
            break
        if _column_has_data_in_rows(data_rows, j):
            inferred.append((j, next_num))
            next_num += 1

    specs = explicit + inferred
    specs.sort(key=lambda x: x[0])
    return specs


def parse_weekly_schedule_excel(file_path: Path, *, lang: str = "ru") -> list[dict[str, Any]]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    if sheet.max_row is None or sheet.max_row < 2:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Файл Excel пуст.", "The Excel file is empty."),
        )
    max_col = int(sheet.max_column or 1)
    last_row = int(sheet.max_row)
    rows_raw = list(
        sheet.iter_rows(
            min_row=1,
            max_row=last_row,
            min_col=1,
            max_col=max_col,
            values_only=True,
        )
    )
    rows = [_pad_row_to_width(r, max_col) for r in rows_raw]
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]

    aliases = {"день недели", "день_недели", "weekday", "day of week"}
    weekday_col_idx: int | None = None
    for i, h in enumerate(headers):
        hl = str(h).strip().lower().replace("ё", "е")
        if hl in aliases:
            weekday_col_idx = i
            break
    if weekday_col_idx is None or _column_index_by_header_aliases(headers, _CLASS_COLUMN_ALIASES) is None:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                lang,
                "Ожидаются колонки «День недели» (или Weekday) и «Класс» / «Группа» (или Class / Group).",
                "Expected columns «День недели» (or Weekday) and «Класс» / «Group» (or Class / Group).",
            ),
        )

    class_col_idx = _column_index_by_header_aliases(headers, _CLASS_COLUMN_ALIASES)
    assert class_col_idx is not None
    data_rows = rows[1:]
    lesson_specs = _discover_lesson_specs(headers, class_col_idx, data_rows)
    if not lesson_specs:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                lang,
                "Нужны колонки «Урок1» / «Слот1» / «Slot1» и т.д.",
                "Slot columns are required (e.g. Урок1, Слот1, Slot1, …).",
            ),
        )

    result: list[dict[str, Any]] = []
    for i, row in enumerate(data_rows):
        excel_row = i + 2
        raw_wd = row[weekday_col_idx] if weekday_col_idx < len(row) else None
        raw_class = row[class_col_idx] if class_col_idx < len(row) else None
        raw_class = _resolve_class_cell_value(sheet, excel_row, class_col_idx, raw_class)
        wd = parse_weekday_value(raw_wd)
        class_label = class_name_from_excel(raw_class)
        if wd is None or not class_label:
            continue

        lessons = []
        for j, lix in lesson_specs:
            value = row[j] if j < len(row) else None
            lessons.append({"index": lix, "subject": "" if value is None else str(value).strip()})

        result.append(
            {
                "weekday": wd,
                "class_name": class_label,
                "class_key": normalize_class(class_label),
                "lessons": lessons,
            }
        )
    return result


def parse_excel(file_path: Path, *, lang: str = "ru") -> list[dict[str, Any]]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    if sheet.max_row is None or sheet.max_row < 2:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Файл Excel пуст.", "The Excel file is empty."),
        )
    max_col = int(sheet.max_column or 1)
    last_row = int(sheet.max_row)
    rows_raw = list(
        sheet.iter_rows(
            min_row=1,
            max_row=last_row,
            min_col=1,
            max_col=max_col,
            values_only=True,
        )
    )
    rows = [_pad_row_to_width(r, max_col) for r in rows_raw]
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    date_col_idx = _column_index_by_header_aliases(headers, _DATE_COLUMN_ALIASES)
    class_col_idx = _column_index_by_header_aliases(headers, _CLASS_COLUMN_ALIASES)
    if date_col_idx is None or class_col_idx is None:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                lang,
                "Ожидаются колонки «Дата» / «Date» и «Класс» / «Группа» (или Class / Group).",
                "Expected columns «Дата» / «Date» and «Класс» / «Group» (or Class / Group).",
            ),
        )

    data_rows = rows[1:]
    lesson_specs = _discover_lesson_specs(headers, class_col_idx, data_rows)
    if not lesson_specs:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                lang,
                "Нужны колонки «Урок1» / «Слот1» / «Slot1» и т.д.",
                "Slot columns are required (e.g. Урок1, Слот1, Slot1, …).",
            ),
        )

    result: list[dict[str, Any]] = []
    for i, row in enumerate(data_rows):
        excel_row = i + 2
        raw_date = row[date_col_idx] if date_col_idx < len(row) else None
        raw_class = row[class_col_idx] if class_col_idx < len(row) else None
        raw_class = _resolve_class_cell_value(sheet, excel_row, class_col_idx, raw_class)
        class_label = class_name_from_excel(raw_class)
        if not raw_date or not class_label:
            continue

        if isinstance(raw_date, datetime):
            parsed_date = raw_date.date()
        elif isinstance(raw_date, date):
            parsed_date = raw_date
        else:
            parsed_date = datetime.strptime(str(raw_date).strip(), "%d.%m.%Y").date()

        lessons = []
        for j, lix in lesson_specs:
            value = row[j] if j < len(row) else None
            lessons.append({"index": lix, "subject": "" if value is None else str(value).strip()})

        result.append(
            {
                "date": parsed_date.isoformat(),
                "class_name": class_label,
                "class_key": normalize_class(class_label),
                "lessons": lessons,
            }
        )
    return result


def parse_holidays_excel(file_path: Path, *, lang: str = "ru") -> list[dict[str, Any]]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Файл праздников Excel пуст.", "The holidays Excel file is empty."),
        )

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    required = {"Дата", "Название", "Описание"}
    if not required.issubset(set(headers)):
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                lang,
                "Ожидаются колонки 'Дата', 'Название', 'Описание'.",
                "Expected columns 'Дата', 'Название', 'Описание'.",
            ),
        )

    result: list[dict[str, Any]] = []
    for row in rows[1:]:
        if row is None:
            continue
        item = dict(zip(headers, row))
        raw_date = item.get("Дата")
        raw_name = item.get("Название")
        raw_description = item.get("Описание")
        if not raw_date or not raw_name:
            continue

        kind: str
        parsed_date: date | None = None
        md: str | None = None
        if isinstance(raw_date, datetime):
            parsed_date = raw_date.date()
            kind = "once"
        elif isinstance(raw_date, date):
            parsed_date = raw_date
            kind = "once"
        else:
            s0 = str(raw_date).strip()
            # допускаем пробелы вокруг точек и хвостовую точку: "01.01", "01. 01", "01.01.", "01 . 01 . 2026"
            s = s0.replace("\u00a0", " ")
            s = re.sub(r"\s+", "", s)
            s = s.rstrip(".")
            try:
                # dd.mm.yyyy -> once, dd.mm -> annual
                if re.fullmatch(r"\d{1,2}\.\d{1,2}\.\d{4}", s):
                    parsed_date = datetime.strptime(s, "%d.%m.%Y").date()
                    kind = "once"
                elif re.fullmatch(r"\d{1,2}\.\d{1,2}", s):
                    d_s, m_s = s.split(".")
                    md = f"{int(m_s):02d}-{int(d_s):02d}"
                    kind = "annual"
                else:
                    # неизвестный формат — пропускаем строку вместо падения ТВ
                    continue
            except (ValueError, TypeError):
                continue

        row_out: dict[str, Any] = {
            "name": str(raw_name).strip(),
            "description": "" if raw_description is None else str(raw_description).strip(),
            "kind": kind,
        }
        if kind == "annual" and md:
            row_out["md"] = md  # MM-DD
        elif parsed_date is not None:
            row_out["date"] = parsed_date.isoformat()
        else:
            continue
        result.append(row_out)
    # стабильная сортировка для хранения: annual по md, once по date
    def _k(item: dict[str, Any]) -> tuple[int, str]:
        if item.get("kind") == "annual":
            return (0, str(item.get("md") or ""))
        return (1, str(item.get("date") or "9999-12-31"))

    result.sort(key=_k)
    return result


def maybe_import_schedule_from_folder() -> None:
    signature = import_signature(AUTO_SCHEDULE_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("schedule")
    if current == signature and SCHEDULE_PATH.exists():
        return
    parsed = parse_excel(AUTO_SCHEDULE_IMPORT_PATH)
    write_json(SCHEDULE_PATH, parsed)
    state["schedule"] = signature
    save_import_state(state)


def maybe_import_full_schedule_from_folder() -> None:
    signature = import_signature(AUTO_FULL_SCHEDULE_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("full_schedule")
    if current == signature and FULL_SCHEDULE_PATH.exists():
        return
    parsed = parse_weekly_schedule_excel(AUTO_FULL_SCHEDULE_IMPORT_PATH)
    write_json(FULL_SCHEDULE_PATH, parsed)
    state["full_schedule"] = signature
    save_import_state(state)


def maybe_import_schedule_sample_from_folder() -> None:
    signature = import_signature(AUTO_SCHEDULE_SAMPLE_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("schedule_sample")
    if current == signature and SCHEDULE_SAMPLE_PATH.exists():
        return
    parsed = parse_weekly_schedule_excel(AUTO_SCHEDULE_SAMPLE_IMPORT_PATH)
    write_json(SCHEDULE_SAMPLE_PATH, parsed)
    state["schedule_sample"] = signature
    save_import_state(state)


def maybe_import_holidays_from_folder() -> None:
    signature = import_signature(AUTO_HOLIDAYS_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("holidays")
    if current == signature and HOLIDAYS_PATH.exists():
        return
    parsed = parse_holidays_excel(AUTO_HOLIDAYS_IMPORT_PATH)
    write_json(HOLIDAYS_PATH, parsed)
    state["holidays"] = signature
    save_import_state(state)


def parse_announcements_excel(file_path: Path, *, lang: str = "ru") -> list[dict[str, Any]]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Файл объявлений Excel пуст.", "The announcements Excel file is empty."),
        )

    # Берём первую колонку; заголовок может быть любым, пропускаем первую строку если похожа на заголовок.
    lines: list[str] = []
    for i, row in enumerate(rows):
        if row is None:
            continue
        v = row[0] if len(row) else None
        if v is None:
            s = ""
        else:
            s = str(v).strip()
        if i == 0 and s.lower() in {"текст", "объявление", "объявления", "announcement", "announcements"}:
            continue
        lines.append(s)

    blocks: list[list[str]] = [[]]
    for raw in lines:
        s = str(raw or "").rstrip()
        if not s:
            # пустые строки сохраняем как пустую строку внутри блока (для форматирования),
            # но не создаём бесконечные пустые блоки
            if blocks and blocks[-1]:
                blocks[-1].append("")
            continue
        st = s.strip()
        if st in ("!!!", "—!!!—"):
            if blocks and blocks[-1]:
                blocks.append([])
            continue
        if st.startswith("//") or st.startswith("#"):
            continue
        blocks[-1].append(s)

    out: list[dict[str, Any]] = []
    for b in blocks:
        text = "\n".join(b).strip("\n")
        if not text.strip():
            continue
        out.append({"text": text})
    return out


def maybe_import_announcements_from_folder() -> None:
    signature = import_signature(AUTO_ANNOUNCEMENTS_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("announcements")
    if current == signature and ANNOUNCEMENTS_PATH.exists():
        return
    parsed = parse_announcements_excel(AUTO_ANNOUNCEMENTS_IMPORT_PATH)
    write_json(ANNOUNCEMENTS_PATH, parsed)
    state["announcements"] = signature
    save_import_state(state)


def parse_marquee_excel(file_path: Path, *, lang: str = "ru") -> list[str]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Файл бегущей строки Excel пуст.", "The marquee Excel file is empty."),
        )
    out: list[str] = []
    for i, row in enumerate(rows):
        if row is None:
            continue
        v = row[0] if len(row) else None
        s = "" if v is None else str(v).strip()
        if i == 0 and s.lower() in {"текст", "бегущая строка", "marquee"}:
            continue
        if not s:
            continue
        if s.startswith("//") or s.startswith("#"):
            continue
        out.append(s)
    return out


def maybe_import_marquee_from_folder() -> None:
    signature = import_signature(AUTO_MARQUEE_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("marquee")
    if current == signature and MARQUEE_PATH.exists():
        return
    parsed = parse_marquee_excel(AUTO_MARQUEE_IMPORT_PATH)
    write_json(MARQUEE_PATH, parsed)
    state["marquee"] = signature
    save_import_state(state)
