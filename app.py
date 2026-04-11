from __future__ import annotations

from contextlib import asynccontextmanager

import hashlib
import hmac
import io
import ipaddress
import json
import re
import secrets
import shutil
import sys
import time
import zipfile
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from typing import Any
from urllib.parse import urlparse

from fastapi import Body, FastAPI, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook


def admin_ui_lang(request: Request) -> str:
    return "en" if request.headers.get("X-UI-Locale", "").strip().lower() == "en" else "ru"


def admin_msg(lang: str, ru: str, en: str) -> str:
    return en if lang == "en" else ru


def session_cookie_secure(request: Request) -> bool:
    xfp = str(request.headers.get("x-forwarded-proto", "")).strip().lower()
    if xfp.startswith("https"):
        return True
    try:
        return request.url.scheme == "https"
    except Exception:
        return False


if getattr(sys, "frozen", False):
    APP_DIR = Path(sys.executable).resolve().parent
    RESOURCE_DIR = Path(getattr(sys, "_MEIPASS", APP_DIR))
else:
    APP_DIR = Path(__file__).resolve().parent
    RESOURCE_DIR = APP_DIR

def resolve_brand_logo_path() -> Path | None:
    for candidate in (APP_DIR / "ico.png", RESOURCE_DIR / "ico.png"):
        if candidate.is_file():
            return candidate
    return None


DATA_DIR = APP_DIR / "data"
UPLOADS_DIR = DATA_DIR / "uploads"
BELL_SOUNDS_DIR = UPLOADS_DIR / "bells"
BACKGROUND_UPLOAD_IMAGE_SUFFIXES = frozenset(
    {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".avif"}
)
BREAK_MUSIC_DIR = DATA_DIR / "break_music"
IMPORT_DIR = DATA_DIR / "import"
STATIC_DIR = RESOURCE_DIR / "static"
CONFIG_PATH = DATA_DIR / "config.json"
AUTH_PATH = DATA_DIR / "auth.json"
SCHEDULE_PATH = DATA_DIR / "schedule.json"
FULL_SCHEDULE_PATH = DATA_DIR / "full_schedule.json"
SCHEDULE_SAMPLE_PATH = DATA_DIR / "schedule_sample.json"
HOLIDAYS_PATH = DATA_DIR / "holidays.json"
ANNOUNCEMENTS_PATH = DATA_DIR / "announcements.json"
MARQUEE_PATH = DATA_DIR / "marquee.json"
OVERRIDES_PATH = DATA_DIR / "overrides.json"
BELL_SCHEDULES_PATH = DATA_DIR / "bell_schedules.json"
CHANGE_LOG_PATH = DATA_DIR / "change_log.json"
APP_VERSION = "1.01.002"
IMPORT_STATE_PATH = DATA_DIR / "import_state.json"
AUTO_SCHEDULE_IMPORT_PATH = IMPORT_DIR / "schedule.xlsx"
AUTO_FULL_SCHEDULE_IMPORT_PATH = IMPORT_DIR / "full_schedule.xlsx"
AUTO_SCHEDULE_SAMPLE_IMPORT_PATH = IMPORT_DIR / "schedule_sample.xlsx"
FULL_SCHEDULE_SAMPLE_XLSX = IMPORT_DIR / "full_schedule_sample.xlsx"
AUTO_HOLIDAYS_IMPORT_PATH = IMPORT_DIR / "holidays.xlsx"
AUTO_ANNOUNCEMENTS_IMPORT_PATH = IMPORT_DIR / "announcements.xlsx"
AUTO_MARQUEE_IMPORT_PATH = IMPORT_DIR / "marquee.xlsx"
SESSION_COOKIE = "gornii_session"
GRID_COLS = 32
GRID_ROWS = 26
WIDGET_IMAGES_SUBDIR = "widget_images"
SINGLETON_WIDGET_IDS = {
    "date": "date",
    "time": "time",
    "text": "text",
    "bell_status": "bell_status",
    "bell_countdown": "bell_countdown",
    "schedule": "schedule",
    "holidays": "holidays",
    "announcements": "announcements",
    "marquee": "marquee",
    "emergency": "emergency",
    "image": "image",
}

DEFAULT_BELL_TRIGGER_SEC_WINDOW = 25
PRE_BELL_LEAD_MINUTES = 1
PRE_BELL_FIRE_SEC_WINDOW = 12
PRE_BELL_MAX_DURATION_SEC = 52
PRE_BELL_AFADE_IN_SEC = 3.0

BREAK_ORCH_HEAD_SEC = 60
BREAK_ORCH_FADE_SEC = 120
BREAK_ORCH_SILENCE_SEC = 60
BREAK_ORCH_START_BELL_SEC = 60
MORNING_PRE_FIRST_LESSON_MIN = 30


def audio_trigger_sec_window(audio: dict[str, Any]) -> int:
    raw = audio.get("bell_trigger_sec_window", DEFAULT_BELL_TRIGGER_SEC_WINDOW)
    if raw is None:
        raw = DEFAULT_BELL_TRIGGER_SEC_WINDOW
    try:
        w = int(raw)
    except (TypeError, ValueError):
        w = DEFAULT_BELL_TRIGGER_SEC_WINDOW
    return max(5, min(55, w))


def _safe_rel_uploads_subdir(raw: Any) -> str:
    s = str(raw or "").strip().replace("\\", "/").strip("/")
    if not s or s in (".",):
        return ""
    parts = [p for p in s.split("/") if p and p not in (".", "..")]
    if not parts:
        return ""
    return "/".join(parts)[:120]


def list_background_images_from_uploads(subdir: str = "") -> list[str]:
    """Публичные URL изображений в data/uploads/<subdir> (bells не используем)."""
    ensure_dirs()
    base = UPLOADS_DIR / _safe_rel_uploads_subdir(subdir)
    if not base.is_dir():
        return []
    names: list[str] = []
    for p in base.iterdir():
        if not p.is_file():
            continue
        if p.name.startswith("."):
            continue
        if p.suffix.lower() not in BACKGROUND_UPLOAD_IMAGE_SUFFIXES:
            continue
        names.append(p.name)
    names.sort(key=str.lower)
    prefix = f"/uploads/{_safe_rel_uploads_subdir(subdir)}/" if _safe_rel_uploads_subdir(subdir) else "/uploads/"
    return [f"{prefix}{n}" for n in names]


def list_background_subdirs_from_uploads() -> list[str]:
    """Список подпапок в uploads, где есть хотя бы 1 картинка. Корень обозначается пустой строкой."""
    ensure_dirs()
    out: set[str] = set()
    # корень
    if any(list_background_images_from_uploads("")):
        out.add("")
    if not UPLOADS_DIR.is_dir():
        return [""]
    for p in UPLOADS_DIR.iterdir():
        if not p.is_dir():
            continue
        if p.name.startswith(".") or p.name.lower() == "bells" or p.name.lower() == WIDGET_IMAGES_SUBDIR:
            continue
        rel = p.name
        if list_background_images_from_uploads(rel):
            out.add(rel)
    # сортировка: корень первым
    rest = sorted([x for x in out if x], key=str.lower)
    return ["", *rest]


def ensure_dirs() -> None:
    DATA_DIR.mkdir(exist_ok=True)
    UPLOADS_DIR.mkdir(exist_ok=True)
    (UPLOADS_DIR / WIDGET_IMAGES_SUBDIR).mkdir(parents=True, exist_ok=True)
    BELL_SOUNDS_DIR.mkdir(parents=True, exist_ok=True)
    BREAK_MUSIC_DIR.mkdir(parents=True, exist_ok=True)
    IMPORT_DIR.mkdir(exist_ok=True)
    ensure_weekly_schedule_template_file()
    ensure_default_change_log()


def write_weekly_schedule_template_excel(path: Path) -> None:
    """Шаблон Excel: «День недели», «Класс», Урок1…Урок8 (как parse_weekly_schedule_excel)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Расписание"
    headers = ["День недели", "Класс"] + [f"Урок{i}" for i in range(1, 9)]
    ws.append(headers)
    for row in (
        ("Понедельник", "5А", "Математика", "Русский язык", "Литература", "Окр. мир", "", "", "", ""),
        ("Вторник", "5А", "История", "Английский язык", "Физическая культура", "", "", "", "", ""),
        ("Среда", "5А", "Математика", "Русский язык", "Технология", "Музыка", "", "", "", ""),
    ):
        ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def ensure_weekly_schedule_template_file() -> None:
    if FULL_SCHEDULE_SAMPLE_XLSX.exists():
        return
    IMPORT_DIR.mkdir(parents=True, exist_ok=True)
    try:
        write_weekly_schedule_template_excel(FULL_SCHEDULE_SAMPLE_XLSX)
    except OSError:
        pass


def export_weekly_schedule_bundle_bytes() -> bytes:
    ensure_dirs()
    buffer = io.BytesIO()
    readme = (
        "Недельное расписание GuardSchool\n"
        "— full_schedule.json — полное по дням недели (как в админке после загрузки Excel).\n"
        "— schedule_sample.json — образец отличий от полного (подсветка на ТВ).\n"
        "— full_schedule_sample.xlsx — шаблон таблицы для Excel (колонки «День недели», «Класс», Урок1…).\n"
        "Импорт: ZIP с теми же именами файлов в разделе «Уроки».\n"
    )
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "full_schedule.json",
            json.dumps(read_json(FULL_SCHEDULE_PATH, []), ensure_ascii=False, indent=2).encode("utf-8"),
        )
        archive.writestr(
            "schedule_sample.json",
            json.dumps(read_json(SCHEDULE_SAMPLE_PATH, []), ensure_ascii=False, indent=2).encode("utf-8"),
        )
        ensure_weekly_schedule_template_file()
        if FULL_SCHEDULE_SAMPLE_XLSX.is_file():
            archive.writestr(FULL_SCHEDULE_SAMPLE_XLSX.name, FULL_SCHEDULE_SAMPLE_XLSX.read_bytes())
        archive.writestr("README.txt", readme.encode("utf-8"))
    return buffer.getvalue()


def import_weekly_schedule_bundle_bytes(raw_bytes: bytes, *, lang: str = "ru") -> None:
    with zipfile.ZipFile(io.BytesIO(raw_bytes), "r") as archive:
        names = archive.namelist()
        if not names:
            raise HTTPException(
                status_code=400,
                detail=admin_msg(lang, "Пустой ZIP-архив.", "Empty ZIP archive."),
            )
        found: set[str] = set()
        for name in names:
            if Path(name).is_absolute() or ".." in Path(name).parts:
                raise HTTPException(
                    status_code=400,
                    detail=admin_msg(lang, "Недопустимый путь в архиве.", "Invalid path in archive."),
                )
            base = Path(name).name
            if base not in ("full_schedule.json", "schedule_sample.json"):
                continue
            raw = archive.read(name)
            data = json.loads(raw.decode("utf-8"))
            if not isinstance(data, list):
                raise HTTPException(
                    status_code=400,
                    detail=admin_msg(
                        lang,
                        f"{base}: ожидается JSON-массив.",
                        f"{base}: expected a JSON array.",
                    ),
                )
            if base == "full_schedule.json":
                write_json(FULL_SCHEDULE_PATH, data)
            else:
                write_json(SCHEDULE_SAMPLE_PATH, data)
            found.add(base)
        if not found:
            raise HTTPException(
                status_code=400,
                detail=admin_msg(
                    lang,
                    "В архиве нет full_schedule.json или schedule_sample.json.",
                    "The archive must contain full_schedule.json and/or schedule_sample.json.",
                ),
            )


def first_bell_sound_path() -> Path | None:
    ensure_dirs()
    if not BELL_SOUNDS_DIR.exists():
        return None
    for p in sorted(BELL_SOUNDS_DIR.iterdir()):
        if p.is_file():
            return p
    return None


def clear_directory(path: Path) -> None:
    if not path.exists():
        return
    for item in path.iterdir():
        if item.is_file():
            item.unlink()
        elif item.is_dir():
            clear_directory(item)
            item.rmdir()


def default_screen(name: str, slug: str) -> dict[str, Any]:
    return {
        "id": secrets.token_hex(4),
        "name": name,
        "slug": slug,
        "ip_note": "",
        "poll_interval_sec": 10,
        "background_image": "",
        "background_rotate_enabled": False,
        "background_rotate_interval_sec": 3600,
        "background_rotate_folder": "",
        "background_force_image": "",
        "background_rotate_cursor": 0,
        "background_rotate_epoch": 0,
        "tv_text_outline_px": 2,
        "tv_text_outline_color": "rgba(0,0,0,0.85)",
        "selected_classes": ["5", "6", "7", "8"],
        "bell_schedule_template": "standard",
        "weekday_bell_templates": {},
        "widgets": [
            {
                "id": "date",
                "type": "date",
                "title": "Дата",
                "enabled": True,
                "x": 0,
                "y": 0,
                "w": 12,
                "h": 2,
                "settings": {"fontSize": 32, "color": "#ffffff", "bold": False},
            },
            {
                "id": "time",
                "type": "time",
                "title": "Время",
                "enabled": True,
                "x": 12,
                "y": 0,
                "w": 8,
                "h": 2,
                "settings": {"fontSize": 42, "color": "#ffffff", "bold": False},
            },
            {
                "id": "text",
                "type": "text",
                "title": "Текст",
                "enabled": True,
                "x": 20,
                "y": 0,
                "w": 12,
                "h": 2,
                "settings": {
                    "text": "Добро пожаловать",
                    "fontSize": 24,
                    "color": "#ffffff",
                    "background": "rgba(0,0,0,0.35)",
                    "bold": False,
                },
            },
            {
                "id": "bell_status",
                "type": "bell_status",
                "title": "Звонки",
                "enabled": True,
                "x": 24,
                "y": 10,
                "w": 8,
                "h": 3,
                "settings": {
                    "fontSize": 18,
                    "titleFontSize": 18,
                    "color": "#ffffff",
                    "background": "rgba(15,23,42,0.55)",
                    "bold": False,
                },
            },
            {
                "id": "schedule",
                "type": "schedule",
                "title": "Расписание",
                "enabled": True,
                "x": 0,
                "y": 2,
                "w": 24,
                "h": 11,
                "settings": {
                    "showTomorrow": True,
                    "highlightColor": "#8b0000",
                    "sampleDiffColor": "#fee2e2",
                    "headerColor": "#1e3a5f",
                    "bodyColor": "rgba(255,255,255,0.9)",
                    "fontSize": 16,
                    "titleFontSize": 20,
                    "classes": ["5", "6", "7", "8"],
                    "bold": False,
                },
            },
            {
                "id": "carousel",
                "type": "carousel",
                "title": "Карусель",
                "enabled": False,
                "x": 0,
                "y": 2,
                "w": 24,
                "h": 11,
                "settings": {
                    "startDelaySec": 0,
                    "animation": "slide",
                    "childWidgetIds": ["schedule", "text"],
                    "childSlideSec": {"schedule": 60, "text": 30},
                },
            },
            {
                "id": "holidays",
                "type": "holidays",
                "title": "События",
                "enabled": False,
                "x": 24,
                "y": 2,
                "w": 8,
                "h": 5,
                "settings": {
                    "fontSize": 16,
                    "titleFontSize": 18,
                    "color": "#ffffff",
                    "background": "rgba(15,23,42,0.55)",
                    "count": 5,
                    "bold": False,
                },
            },
            {
                "id": "announcements",
                "type": "announcements",
                "title": "Объявления",
                "enabled": False,
                "x": 20,
                "y": 0,
                "w": 12,
                "h": 4,
                "settings": {
                    "items": "Собрание педагогов в 15:00\nЛинейка в понедельник\nПроверить сменную обувь",
                    "fontSize": 18,
                    "titleFontSize": 18,
                    "color": "#ffffff",
                    "background": "rgba(15,23,42,0.55)",
                    "bold": False,
                },
            },
            {
                "id": "marquee",
                "type": "marquee",
                "title": "Бегущая строка",
                "enabled": False,
                "x": 0,
                "y": 24,
                "w": 32,
                "h": 2,
                "settings": {
                    "items": "Внимание! Идет настройка экрана\nСегодня педсовет в 15:00",
                    "fontSize": 22,
                    "color": "#ffffff",
                    "background": "rgba(15,23,42,0.7)",
                    "speedSec": 18,
                    "bold": False,
                },
            },
            {
                "id": "emergency",
                "type": "emergency",
                "title": "Аварийный",
                "enabled": False,
                "x": 0,
                "y": 0,
                "w": GRID_COLS,
                "h": GRID_ROWS,
                "settings": {
                    "text": "ВНИМАНИЕ!\nСрочное сообщение.",
                    "fontSize": 42,
                    "color": "#ffffff",
                    "background": "#b91c1c",
                    "bold": True,
                    "backdrop": False,
                },
            },
            {
                "id": "image",
                "type": "image",
                "title": "Изображение",
                "enabled": False,
                "x": 0,
                "y": 0,
                "w": 8,
                "h": 8,
                "settings": {
                    "images": [{"name": "Эмблема", "url": ""}],
                    "imagesRotateSec": 0,
                    "opacity": 85,
                    "objectFit": "contain",
                    "backdrop": False,
                },
            },
        ],
        "template": "default_schedule",
        "is_active": True,
    }


def normalize_multicast_ip_value(raw: str | None, *, fallback: str) -> str:
    """Если в поле вставили udp://224.x...:port — вытаскиваем IPv4 группы."""
    s = (raw or "").strip()
    if not s:
        return fallback
    try:
        ip = ipaddress.ip_address(s)
        if ip.version == 4:
            return str(ip)
    except ValueError:
        pass
    m = re.search(
        r"\b(22[4-9]|23[0-9])(?:\.(?:25[0-5]|2[0-4]\d|[01]?\d?\d)){3}\b",
        s,
    )
    if m:
        try:
            ip = ipaddress.ip_address(m.group(0))
            if ip.version == 4 and ip.is_multicast:
                return str(ip)
        except ValueError:
            pass
    return fallback


def default_audio_stream_settings() -> dict[str, Any]:
    """Настройки звука на ПК (ffplay): звонки по расписанию и фон на переменах. Поля сети в JSON могут остаться от старой версии — не используются."""
    return {
        "enabled": False,
        "receiver_ip": "",
        "stream_port": 11990,
        "ping_host": "",
        "multicast_ip": "224.0.224.1",
        "multicast_ttl": 10,
        "base_port": 11990,
        "ffmpeg_path": "",
        "interface_note": "",
        "use_bell_schedule": True,
        "use_bell_sound_files": True,
        "volume_percent": 80,
        "source_screen_id": "",
        "send_via_multicast": False,
        "udp_bind_localaddr": False,
        "stream_profile": "mpegts_aac",
        "break_music_on_breaks": False,
        "break_music_volume_percent": 40,
        "bell_trigger_sec_window": DEFAULT_BELL_TRIGGER_SEC_WINDOW,
    }


def sanitize_audio_stream(raw: dict[str, Any] | None) -> dict[str, Any]:
    d = default_audio_stream_settings()
    if not isinstance(raw, dict):
        return d
    d["enabled"] = bool(raw.get("enabled"))
    recv = normalize_stream_host(str(raw.get("receiver_ip") or raw.get("ping_host") or ""))
    d["receiver_ip"] = recv
    d["ping_host"] = recv
    d["multicast_ip"] = normalize_multicast_ip_value(
        str(raw.get("multicast_ip") or ""),
        fallback=d["multicast_ip"],
    )[:45]
    d["ffmpeg_path"] = str(raw.get("ffmpeg_path") or "").strip()[:500]
    d["interface_note"] = str(raw.get("interface_note") or "").strip()[:500]
    try:
        ttl = int(raw.get("multicast_ttl", d["multicast_ttl"]))
        d["multicast_ttl"] = max(1, min(255, ttl))
    except (TypeError, ValueError):
        pass
    try:
        port = int(raw.get("stream_port", raw.get("base_port", d["stream_port"])))
        port = max(1, min(65535, port))
        d["stream_port"] = port
        d["base_port"] = port
    except (TypeError, ValueError):
        pass
    d["use_bell_schedule"] = bool(raw.get("use_bell_schedule", True))
    d["use_bell_sound_files"] = bool(raw.get("use_bell_sound_files", True))
    try:
        vol = int(raw.get("volume_percent", d["volume_percent"]))
        d["volume_percent"] = max(0, min(100, vol))
    except (TypeError, ValueError):
        pass
    try:
        ipaddress.ip_address(d["multicast_ip"])
    except ValueError:
        d["multicast_ip"] = default_audio_stream_settings()["multicast_ip"]
    if d["ffmpeg_path"]:
        p = Path(d["ffmpeg_path"])
        if not p.is_file() and not shutil.which(d["ffmpeg_path"]):
            d["ffmpeg_path"] = ""
    d["source_screen_id"] = str(raw.get("source_screen_id") or "").strip()[:80]
    d["send_via_multicast"] = bool(raw.get("send_via_multicast"))
    d["udp_bind_localaddr"] = bool(raw.get("udp_bind_localaddr"))
    prof = str(raw.get("stream_profile") or d["stream_profile"]).strip().lower()
    if prof not in ("mpegts_aac", "mpegts_copy", "mpegts_mp3"):
        prof = "mpegts_aac"
    d["stream_profile"] = prof
    d["break_music_on_breaks"] = bool(raw.get("break_music_on_breaks"))
    try:
        bv = int(raw.get("break_music_volume_percent", d["break_music_volume_percent"]))
        d["break_music_volume_percent"] = max(0, min(100, bv))
    except (TypeError, ValueError):
        pass
    try:
        bw = int(raw.get("bell_trigger_sec_window", d["bell_trigger_sec_window"]))
        d["bell_trigger_sec_window"] = max(5, min(55, bw))
    except (TypeError, ValueError):
        pass
    return d


def normalize_stream_host(raw: str) -> str:
    """
    Из поля «IP приёмника» часто вставляют URL (http://192.168.1.1/) — ping и TCP ждут голый хост.
    Убираем схему, путь, лишние слэши; при виде 192.168.1.1:порт оставляем только IP (порт задаётся отдельно).
    """
    s = str(raw or "").strip()
    if not s:
        return ""
    if re.match(r"^https?://", s, re.IGNORECASE):
        try:
            p = urlparse(s)
            h = (p.hostname or "").strip()
            if h:
                return h[:253]
        except Exception:
            pass
    s = s.split("/")[0].split("?")[0].strip()
    if s.count(":") == 1 and not s.startswith("["):
        left, _, right = s.partition(":")
        if right.isdigit():
            s = left.strip()
    return s[:253]


def default_config() -> dict[str, Any]:
    return {
        "templateSystem": {"version": 2, "templates": ["default_schedule"], "grid": {"cols": GRID_COLS, "rows": GRID_ROWS}},
        "ui_locale": "ru",
        "timezone": "Europe/Moscow",
        "clock_offset_minutes": 0,
        "screens": [default_screen("ТВ-1", "tv-1")],
        "audio_stream": default_audio_stream_settings(),
    }


def _valid_iana_timezone(name: str) -> bool:
    s = (name or "").strip()
    if not s or len(s) > 80:
        return False
    try:
        ZoneInfo(s)
        return True
    except Exception:
        return False


def calendar_today_for_config(config: dict[str, Any]) -> date:
    """«Сегодня» для расписания и звонков по часовому поясу из конфига."""
    tzname = str((config or {}).get("timezone") or "Europe/Moscow").strip()
    try:
        z = ZoneInfo(tzname)
    except Exception:
        z = ZoneInfo("Europe/Moscow")
    return datetime.now(z).date()


def display_settings_dict(config: dict[str, Any]) -> dict[str, Any]:
    return {
        "timezone": str((config or {}).get("timezone") or "Europe/Moscow"),
        "clock_offset_minutes": int((config or {}).get("clock_offset_minutes") or 0),
        "ui_locale": str((config or {}).get("ui_locale") or "ru"),
    }


def default_bell_schedules() -> dict[str, Any]:
    return {
        "templates": [
            {
                "id": "standard",
                "name": "Обычное",
                "entries": [
                    {"lesson": "1", "start": "08:30", "end": "09:15"},
                    {"lesson": "2", "start": "09:25", "end": "10:10"},
                    {"lesson": "3", "start": "10:30", "end": "11:15"},
                    {"lesson": "4", "start": "11:35", "end": "12:20"},
                    {"lesson": "5", "start": "12:30", "end": "13:15"},
                    {"lesson": "6", "start": "13:25", "end": "14:10"},
                    {"lesson": "7", "start": "14:20", "end": "15:05"},
                    {"lesson": "8", "start": "15:15", "end": "16:00"},
                ],
            },
            {
                "id": "short",
                "name": "Сокращенное",
                "entries": [
                    {"lesson": "1", "start": "08:30", "end": "09:05"},
                    {"lesson": "2", "start": "09:15", "end": "09:50"},
                    {"lesson": "3", "start": "10:00", "end": "10:35"},
                    {"lesson": "4", "start": "10:45", "end": "11:20"},
                    {"lesson": "5", "start": "11:30", "end": "12:05"},
                    {"lesson": "6", "start": "12:15", "end": "12:50"},
                    {"lesson": "7", "start": "13:00", "end": "13:35"},
                    {"lesson": "8", "start": "13:45", "end": "14:20"},
                ],
            },
        ],
        "date_overrides": [],
        "weekday_overrides": {},
        "sound_defaults": {"start": None, "end": None},
    }


def read_json(path: Path, fallback: Any) -> Any:
    if not path.exists():
        return fallback
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def load_change_log() -> list[dict[str, Any]]:
    raw = read_json(CHANGE_LOG_PATH, [])
    if not isinstance(raw, list):
        return []
    out: list[dict[str, Any]] = []
    for item in raw:
        if not isinstance(item, dict):
            continue
        msg = item.get("message")
        if msg is None:
            continue
        ts = item.get("timestamp")
        if not ts:
            ts = datetime.now().isoformat(timespec="seconds")
        ver = item.get("version")
        out.append(
            {
                "timestamp": str(ts),
                "message": str(msg),
                "version": "" if ver is None else str(ver).strip(),
            }
        )
    return out


def ensure_default_change_log() -> None:
    """Одна стартовая запись для новой установки (если файла нет или он пустой)."""
    if CHANGE_LOG_PATH.exists():
        try:
            raw = json.loads(CHANGE_LOG_PATH.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            raw = []
        if isinstance(raw, list) and len(raw) > 0:
            return
    write_json(
        CHANGE_LOG_PATH,
        [
            {
                "version": APP_VERSION,
                "timestamp": datetime.now().isoformat(timespec="seconds"),
                "message": (
                    "Первая установка: журнал выпусков в data/change_log.json. При каждом релизе поднимайте APP_VERSION "
                    "и добавляйте сюда краткое описание изменений."
                ),
            },
        ],
    )


def append_release_note(message: str, *, version: str | None = None) -> None:
    """Добавить запись о версии при выпуске (поднять APP_VERSION и вызвать из кода или править JSON вручную)."""
    v = (version or APP_VERSION).strip()
    entries = load_change_log()
    entries.insert(0, {"version": v, "timestamp": datetime.now().isoformat(timespec="seconds"), "message": message.strip()})
    write_json(CHANGE_LOG_PATH, entries[:200])


def load_import_state() -> dict[str, Any]:
    return read_json(IMPORT_STATE_PATH, {})


def save_import_state(state: dict[str, Any]) -> None:
    write_json(IMPORT_STATE_PATH, state)


def import_signature(path: Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    stat = path.stat()
    return {"mtime_ns": stat.st_mtime_ns, "size": stat.st_size}


def export_bundle_bytes() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in [
            CONFIG_PATH,
            AUTH_PATH,
            SCHEDULE_PATH,
            FULL_SCHEDULE_PATH,
            SCHEDULE_SAMPLE_PATH,
            OVERRIDES_PATH,
            BELL_SCHEDULES_PATH,
            CHANGE_LOG_PATH,
        ]:
            if path.exists():
                archive.writestr(path.name, path.read_bytes())
        if UPLOADS_DIR.exists():
            for file_path in UPLOADS_DIR.rglob("*"):
                if file_path.is_file():
                    archive.writestr(str(Path("uploads") / file_path.relative_to(UPLOADS_DIR)), file_path.read_bytes())
    return buffer.getvalue()


def import_bundle_bytes(raw_bytes: bytes, *, lang: str = "ru") -> None:
    with zipfile.ZipFile(io.BytesIO(raw_bytes), "r") as archive:
        names = archive.namelist()
        if not names:
            raise HTTPException(
                status_code=400,
                detail=admin_msg(lang, "Пустой ZIP-архив.", "Empty ZIP archive."),
            )
        allowed_files = {
            "config.json",
            "auth.json",
            "schedule.json",
            "full_schedule.json",
            "schedule_sample.json",
            "overrides.json",
            "bell_schedules.json",
            "change_log.json",
        }
        for name in names:
            path = Path(name)
            if path.is_absolute() or ".." in path.parts:
                raise HTTPException(
                    status_code=400,
                    detail=admin_msg(lang, "Архив содержит недопустимые пути.", "The archive contains invalid paths."),
                )
            if path.parts and path.parts[0] == "uploads":
                continue
            if path.name not in allowed_files or len(path.parts) != 1:
                raise HTTPException(
                    status_code=400,
                    detail=admin_msg(
                        lang,
                        f"Неизвестный файл в архиве: {name}",
                        f"Unknown file in archive: {name}",
                    ),
                )

        has_root_config = any(
            (not Path(n).is_absolute() and ".." not in Path(n).parts and len(Path(n).parts) == 1 and Path(n).name == "config.json")
            for n in names
        )
        if not has_root_config:
            raise HTTPException(
                status_code=400,
                detail=admin_msg(
                    lang,
                    "В корне ZIP должен быть config.json (экспорт GuardSchool). Импорт отменён, каталог uploads не очищался.",
                    "ZIP must contain config.json at archive root (GuardSchool export). Import aborted; uploads were not cleared.",
                ),
            )

        clear_directory(UPLOADS_DIR)
        ensure_dirs()

        for name in names:
            path = Path(name)
            payload = archive.read(name)
            if path.parts and path.parts[0] == "uploads":
                target = UPLOADS_DIR / Path(*path.parts[1:])
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(payload)
            else:
                (DATA_DIR / path.name).write_bytes(payload)


def normalize_widget(widget: dict[str, Any]) -> dict[str, Any]:
    widget["id"] = widget.get("id") or SINGLETON_WIDGET_IDS.get(widget.get("type"), secrets.token_hex(4))
    widget.setdefault("enabled", True)
    widget.setdefault("settings", {})
    if widget["type"] in {"date", "time", "text", "bell_status", "bell_countdown"}:
        widget["settings"].setdefault("fontSize", 24)
        widget["settings"].setdefault("color", "#ffffff")
        widget["settings"].setdefault("bold", False)
    if widget["type"] in {"holidays", "announcements", "marquee"}:
        widget["settings"].setdefault("fontSize", 18)
        widget["settings"].setdefault("color", "#ffffff")
        widget["settings"].setdefault("background", "rgba(15,23,42,0.55)")
        widget["settings"].setdefault("bold", False)
    if widget["type"] in {"holidays", "announcements"}:
        widget["settings"].setdefault("titleFontSize", 18)
    if widget["type"] == "text":
        widget["settings"].setdefault("background", "rgba(0,0,0,0.35)")
        widget["settings"].setdefault("text", "")
    if widget["type"] == "announcements":
        widget["settings"].setdefault("items", "")
        widget["settings"].setdefault("useManual", False)
        widget["settings"].setdefault("rotateSec", 30)
    if widget["type"] == "marquee":
        widget["settings"].setdefault("items", "")
        widget["settings"].setdefault("useManual", False)
        widget["settings"].setdefault("speedSec", 18)
    if widget["type"] == "holidays":
        widget["settings"].setdefault("count", 5)
    if widget["type"] in {"bell_status", "bell_countdown"}:
        widget["settings"].setdefault("background", "rgba(15,23,42,0.55)")
        widget["settings"].setdefault("titleFontSize", 18)
    if widget["type"] == "schedule":
        widget["settings"].setdefault("highlightColor", "#8b0000")
        widget["settings"].setdefault("sampleDiffColor", "#fee2e2")
        widget["settings"].setdefault("headerColor", "#1e3a5f")
        widget["settings"].setdefault("bodyColor", "rgba(255,255,255,0.9)")
        widget["settings"].setdefault("fontSize", 16)
        widget["settings"].setdefault("titleFontSize", 20)
        widget["settings"].setdefault("classes", [])
        widget["settings"].setdefault("bold", False)
    if widget["type"] == "carousel":
        widget["settings"].setdefault("startDelaySec", 0)
        widget["settings"].setdefault("animation", "slide")
        widget["settings"].setdefault("childWidgetIds", [])
        widget["settings"].setdefault("childSlideSec", {})
    if widget["type"] == "emergency":
        widget["x"] = 0
        widget["y"] = 0
        widget["w"] = GRID_COLS
        widget["h"] = GRID_ROWS
        widget["settings"].setdefault("text", "")
        widget["settings"].setdefault("fontSize", 42)
        widget["settings"].setdefault("color", "#ffffff")
        widget["settings"].setdefault("background", "#b91c1c")
        widget["settings"].setdefault("bold", True)
        widget["settings"].setdefault("backdrop", False)
    if widget["type"] == "image":
        try:
            op = int(widget["settings"].get("opacity", 85))
        except (TypeError, ValueError):
            op = 85
        widget["settings"]["opacity"] = max(0, min(100, op))
        widget["settings"].setdefault("objectFit", "contain")
        widget["settings"].setdefault("backdrop", False)
        legacy_url = str(widget["settings"].get("imageUrl") or "").strip()
        raw_images = widget["settings"].get("images")
        images: list[dict[str, str]] = []
        if isinstance(raw_images, list):
            for item in raw_images:
                if not isinstance(item, dict):
                    continue
                nm = str(item.get("name") or "").strip()[:120]
                u = str(item.get("url") or "").strip()[:512]
                if not nm and not u:
                    continue
                images.append({"name": nm or "Без названия", "url": u})
        if not images and legacy_url:
            images = [{"name": "Изображение", "url": legacy_url[:512]}]
        if not images:
            images = [{"name": "Картинка 1", "url": ""}]
        widget["settings"]["images"] = images
        widget["settings"].pop("imageUrl", None)
        try:
            irs = int(widget["settings"].get("imagesRotateSec", 0))
        except (TypeError, ValueError):
            irs = 0
        widget["settings"]["imagesRotateSec"] = max(0, min(600, irs))
    widget["settings"].setdefault("backdrop", True)
    return widget


def dedupe_widgets(screen: dict[str, Any]) -> dict[str, Any]:
    seen_ids: set[str] = set()
    seen_singletons: set[str] = set()
    result: list[dict[str, Any]] = []
    for raw_widget in screen.get("widgets", []):
        widget = normalize_widget(raw_widget)
        widget_type = widget.get("type")
        widget_id = widget.get("id")
        if widget_type in SINGLETON_WIDGET_IDS:
            if widget_type in seen_singletons:
                continue
            seen_singletons.add(widget_type)
            widget["id"] = SINGLETON_WIDGET_IDS[widget_type]
            widget_id = widget["id"]
        if widget_id in seen_ids:
            continue
        seen_ids.add(widget_id)
        result.append(widget)
    screen["widgets"] = result
    return screen


def ensure_default_widgets(screen: dict[str, Any]) -> dict[str, Any]:
    existing_types = {widget.get("type") for widget in screen.get("widgets", [])}
    for template_widget in default_screen(screen.get("name", "ТВ"), screen.get("slug", "tv"))["widgets"]:
        widget_type = template_widget.get("type")
        if widget_type in SINGLETON_WIDGET_IDS and widget_type not in existing_types:
            screen.setdefault("widgets", []).append(normalize_widget(dict(template_widget)))
    return screen


def migrate_screen_layout(screen: dict[str, Any], old_cols: int, old_rows: int) -> dict[str, Any]:
    if old_cols == GRID_COLS and old_rows == GRID_ROWS:
        return screen
    default_by_type = {widget["type"]: widget for widget in default_screen(screen["name"], screen["slug"])["widgets"]}
    migrated_widgets = []
    for widget in screen.get("widgets", []):
        current = normalize_widget(widget)
        template = default_by_type.get(current["type"])
        if template:
            current["x"] = template["x"]
            current["y"] = template["y"]
            current["w"] = template["w"]
            current["h"] = template["h"]
            for key, value in template["settings"].items():
                current["settings"].setdefault(key, value)
        migrated_widgets.append(current)
    screen["widgets"] = migrated_widgets
    return screen


def load_config() -> dict[str, Any]:
    config = read_json(CONFIG_PATH, default_config())
    if not config.get("screens"):
        config["screens"] = [default_screen("ТВ-1", "tv-1")]
    grid = config.setdefault("templateSystem", {}).setdefault("grid", {"cols": GRID_COLS, "rows": GRID_ROWS})
    old_cols = int(grid.get("cols", GRID_COLS))
    old_rows = int(grid.get("rows", GRID_ROWS))
    for screen in config["screens"]:
        migrate_screen_layout(screen, old_cols, old_rows)
        dedupe_widgets(screen)
        ensure_default_widgets(screen)
        dedupe_widgets(screen)
        screen.setdefault("selected_classes", ["5", "6", "7", "8"])
        screen.setdefault("background_image", "")
        screen.setdefault("background_rotate_enabled", False)
        screen.setdefault("background_rotate_interval_sec", 3600)
        screen.setdefault("background_rotate_folder", "")
        screen.setdefault("background_force_image", "")
        screen.setdefault("background_rotate_cursor", 0)
        screen.setdefault("background_rotate_epoch", 0)
        screen.setdefault("tv_text_outline_px", 2)
        screen.setdefault("tv_text_outline_color", "rgba(0,0,0,0.85)")
        try:
            _bri = int(screen["background_rotate_interval_sec"])
        except (TypeError, ValueError):
            _bri = 3600
        screen["background_rotate_interval_sec"] = max(60, min(86400, _bri))
        screen["background_rotate_enabled"] = bool(screen.get("background_rotate_enabled"))
        screen["background_rotate_folder"] = _safe_rel_uploads_subdir(screen.get("background_rotate_folder"))
        bfi = str(screen.get("background_force_image") or "").strip()
        # Разрешаем только /uploads/... (иначе можно подсунуть внешний URL)
        screen["background_force_image"] = bfi if bfi.startswith("/uploads/") else ""
        try:
            cur = int(screen.get("background_rotate_cursor", 0))
        except (TypeError, ValueError):
            cur = 0
        screen["background_rotate_cursor"] = max(0, min(1000000, cur))
        try:
            ep = int(screen.get("background_rotate_epoch", 0))
        except (TypeError, ValueError):
            ep = 0
        screen["background_rotate_epoch"] = max(0, min(4102444800, ep))  # ~2100
        try:
            _px = int(screen.get("tv_text_outline_px", 2))
        except (TypeError, ValueError):
            _px = 2
        screen["tv_text_outline_px"] = max(0, min(8, _px))
        col = screen.get("tv_text_outline_color")
        screen["tv_text_outline_color"] = str(col).strip()[:64] if col is not None else "rgba(0,0,0,0.85)"
        screen.setdefault("poll_interval_sec", 10)
        screen.setdefault("bell_schedule_template", "standard")
        screen.setdefault("weekday_bell_templates", {})
    config["templateSystem"]["version"] = 2
    config["templateSystem"]["grid"] = {"cols": GRID_COLS, "rows": GRID_ROWS}
    config["audio_stream"] = sanitize_audio_stream(config.get("audio_stream"))
    config.setdefault("ui_locale", "ru")
    config.setdefault("timezone", "Europe/Moscow")
    config.setdefault("clock_offset_minutes", 0)
    if str(config.get("ui_locale") or "ru").strip().lower() not in ("ru", "en"):
        config["ui_locale"] = "ru"
    tz0 = str(config.get("timezone") or "Europe/Moscow").strip()
    config["timezone"] = tz0 if _valid_iana_timezone(tz0) else "Europe/Moscow"
    try:
        _off0 = int(config.get("clock_offset_minutes", 0))
    except (TypeError, ValueError):
        _off0 = 0
    config["clock_offset_minutes"] = max(-720, min(720, _off0))
    return config


def sanitize_config(config: dict[str, Any]) -> dict[str, Any]:
    config = config or default_config()
    template_system = config.setdefault("templateSystem", {})
    template_system["version"] = 2
    template_system["grid"] = {"cols": GRID_COLS, "rows": GRID_ROWS}

    screens = config.get("screens") or [default_screen("ТВ-1", "tv-1")]
    cleaned_screens: list[dict[str, Any]] = []
    for raw_screen in screens:
        screen = dict(raw_screen)
        screen.setdefault("selected_classes", ["5", "6", "7", "8"])
        screen["selected_classes"] = list(dict.fromkeys(screen.get("selected_classes", [])))
        screen.setdefault("background_image", "")
        screen.setdefault("background_rotate_enabled", False)
        screen.setdefault("background_rotate_interval_sec", 3600)
        screen.setdefault("background_rotate_folder", "")
        screen.setdefault("background_force_image", "")
        screen.setdefault("background_rotate_cursor", 0)
        screen.setdefault("background_rotate_epoch", 0)
        screen.setdefault("tv_text_outline_px", 2)
        screen.setdefault("tv_text_outline_color", "rgba(0,0,0,0.85)")
        try:
            _bri2 = int(screen["background_rotate_interval_sec"])
        except (TypeError, ValueError):
            _bri2 = 3600
        screen["background_rotate_interval_sec"] = max(60, min(86400, _bri2))
        screen["background_rotate_enabled"] = bool(screen.get("background_rotate_enabled"))
        screen["background_rotate_folder"] = _safe_rel_uploads_subdir(screen.get("background_rotate_folder"))
        bfi2 = str(screen.get("background_force_image") or "").strip()
        screen["background_force_image"] = bfi2 if bfi2.startswith("/uploads/") else ""
        try:
            cur2 = int(screen.get("background_rotate_cursor", 0))
        except (TypeError, ValueError):
            cur2 = 0
        screen["background_rotate_cursor"] = max(0, min(1000000, cur2))
        try:
            ep2 = int(screen.get("background_rotate_epoch", 0))
        except (TypeError, ValueError):
            ep2 = 0
        screen["background_rotate_epoch"] = max(0, min(4102444800, ep2))
        try:
            _px2 = int(screen.get("tv_text_outline_px", 2))
        except (TypeError, ValueError):
            _px2 = 2
        screen["tv_text_outline_px"] = max(0, min(8, _px2))
        col2 = screen.get("tv_text_outline_color")
        screen["tv_text_outline_color"] = str(col2).strip()[:64] if col2 is not None else "rgba(0,0,0,0.85)"
        screen.setdefault("poll_interval_sec", 10)
        screen.setdefault("bell_schedule_template", "standard")
        screen.setdefault("weekday_bell_templates", {})
        dedupe_widgets(screen)
        ensure_default_widgets(screen)
        dedupe_widgets(screen)

        valid_widget_ids = {widget["id"] for widget in screen["widgets"]}
        for widget in screen["widgets"]:
            if widget["type"] == "schedule":
                classes = widget["settings"].get("classes") or screen["selected_classes"]
                widget["settings"]["classes"] = list(dict.fromkeys(classes))
            if widget["type"] == "carousel":
                child_ids = widget["settings"].get("childWidgetIds", [])
                child_ids = [
                    item
                    for item in child_ids
                    if (item in valid_widget_ids or item == "__blank__") and item != widget["id"]
                ]
                widget["settings"]["childWidgetIds"] = list(dict.fromkeys(child_ids))
                cs = widget["settings"].get("childSlideSec") or {}
                if isinstance(cs, dict):
                    widget["settings"]["childSlideSec"] = {
                        k: v for k, v in cs.items() if (k in valid_widget_ids or k == "__blank__") and k != widget["id"]
                    }
            ws = widget.get("settings")
            if isinstance(ws, dict) and "backdrop" in ws:
                b = ws["backdrop"]
                if isinstance(b, str):
                    ws["backdrop"] = b.strip().lower() in ("true", "1", "yes", "on")
                elif not isinstance(b, bool):
                    ws["backdrop"] = bool(b)

        cleaned_screens.append(screen)

    config["screens"] = cleaned_screens
    config["audio_stream"] = sanitize_audio_stream(config.get("audio_stream"))
    loc = str(config.get("ui_locale") or "ru").strip().lower()
    config["ui_locale"] = loc if loc in ("ru", "en") else "ru"
    tz = str(config.get("timezone") or "Europe/Moscow").strip()
    config["timezone"] = tz if _valid_iana_timezone(tz) else "Europe/Moscow"
    try:
        off = int(config.get("clock_offset_minutes", 0))
    except (TypeError, ValueError):
        off = 0
    config["clock_offset_minutes"] = max(-720, min(720, off))
    return config


def load_auth() -> dict[str, Any]:
    return read_json(AUTH_PATH, {})


def hash_password(password: str, salt: str) -> str:
    raw = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 150000)
    return raw.hex()


def password_is_valid(password: str) -> bool:
    return len(password) >= 8 and any(ch.isalpha() for ch in password) and any(ch.isdigit() for ch in password)


def create_session_token(username: str, auth: dict[str, Any]) -> str:
    signature = hmac.new(
        auth["password_hash"].encode("utf-8"),
        username.encode("utf-8"),
        hashlib.sha256,
    ).hexdigest()
    return f"{username}:{signature}"


def verify_session_token(token: str | None, auth: dict[str, Any]) -> bool:
    if not token or ":" not in token or not auth:
        return False
    username, signature = token.split(":", 1)
    if username != auth.get("username"):
        return False
    expected = create_session_token(username, auth).split(":", 1)[1]
    return hmac.compare_digest(signature, expected)


def is_authenticated(request: Request) -> bool:
    token = request.cookies.get(SESSION_COOKIE)
    return verify_session_token(token, load_auth())


def require_auth(request: Request) -> None:
    loc = admin_ui_lang(request)
    if not load_auth():
        raise HTTPException(
            status_code=428,
            detail=admin_msg(
                loc,
                "Требуется первичная настройка администратора.",
                "Complete initial administrator setup first.",
            ),
        )
    if not is_authenticated(request):
        raise HTTPException(
            status_code=401,
            detail=admin_msg(loc, "Требуется вход.", "Sign in required."),
        )


def load_schedule() -> list[dict[str, Any]]:
    maybe_import_schedule_from_folder()
    return read_json(SCHEDULE_PATH, [])


def load_full_schedule() -> list[dict[str, Any]]:
    maybe_import_full_schedule_from_folder()
    return read_json(FULL_SCHEDULE_PATH, [])


def load_schedule_sample() -> list[dict[str, Any]]:
    maybe_import_schedule_sample_from_folder()
    return read_json(SCHEDULE_SAMPLE_PATH, [])


def load_holidays() -> list[dict[str, Any]]:
    maybe_import_holidays_from_folder()
    return read_json(HOLIDAYS_PATH, [])


def load_announcements() -> list[dict[str, Any]]:
    maybe_import_announcements_from_folder()
    return read_json(ANNOUNCEMENTS_PATH, [])


def load_marquee_items() -> list[str]:
    maybe_import_marquee_from_folder()
    raw = read_json(MARQUEE_PATH, [])
    return [str(x).strip() for x in raw if str(x).strip()]


def load_overrides() -> list[dict[str, Any]]:
    return read_json(OVERRIDES_PATH, [])


def migrate_bell_schedules(raw: dict[str, Any]) -> dict[str, Any]:
    data = dict(raw)
    data.setdefault("sound_defaults", {"start": None, "end": None})
    sd = data.get("sound_defaults")
    if not isinstance(sd, dict):
        data["sound_defaults"] = {"start": None, "end": None}
    else:
        sd.setdefault("start", None)
        sd.setdefault("end", None)

    def fix_entries(entries: list[dict[str, Any]]) -> None:
        for ent in entries:
            if ent.get("is_lesson") is False and str(ent.get("label") or "").strip():
                if not str(ent.get("lesson") or "").strip():
                    ent["lesson"] = str(ent["label"]).strip()
            if "lesson" in ent and ent["lesson"] is not None and not isinstance(ent["lesson"], str):
                ent["lesson"] = str(ent["lesson"]).strip()

    for tpl in data.get("templates", []):
        tpl.setdefault("last_lesson", None)
        fix_entries(tpl.get("entries", []))
    for ov in data.get("date_overrides", []):
        ov.setdefault("last_lesson", None)
        fix_entries(ov.get("entries", []))
    return data


def load_bell_schedules() -> dict[str, Any]:
    return migrate_bell_schedules(read_json(BELL_SCHEDULES_PATH, default_bell_schedules()))


def normalize_class(value: str) -> str:
    return str(value).strip().lower()


def class_matches_selector(class_key: str, selector: str) -> bool:
    normalized_selector = normalize_class(selector)
    if not normalized_selector:
        return False
    if class_key == normalized_selector:
        return True
    grade_match = re.match(r"^(\d+)$", normalized_selector)
    if grade_match:
        return re.match(rf"^{grade_match.group(1)}\s*[a-zа-я]+$", class_key) is not None
    return False


def class_in_selected(class_key: str, selectors: list[str]) -> bool:
    if not selectors:
        return True
    return any(class_matches_selector(class_key, selector) for selector in selectors)


def class_sort_key(value: str) -> tuple[int, str]:
    normalized = normalize_class(value)
    match = re.match(r"^(\d+)\s*([a-zа-я]*)$", normalized)
    if match:
        return int(match.group(1)), match.group(2)
    return 10_000, normalized


def time_to_minutes(value: str) -> int:
    hours, minutes = value.split(":")
    return int(hours) * 60 + int(minutes)


def minutes_phrase(n: int) -> str:
    """Склонение для фразы «N минут» в статусе звонков."""
    n = int(n)
    if n % 10 == 1 and n % 100 != 11:
        return f"{n} минута"
    if n % 10 in (2, 3, 4) and n % 100 not in (12, 13, 14):
        return f"{n} минуты"
    return f"{n} минут"


def as_lesson_number(value: Any) -> int | None:
    """Номер урока в шаблоне / расписании может прийти строкой из JSON или Excel."""
    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def entry_academic_num(entry: dict[str, Any]) -> int | None:
    """Только если в «Урок» одни цифры — номер для связи с расписанием Excel (Урок1…8)."""
    raw = entry.get("lesson")
    if raw is None:
        return None
    s = str(raw).strip()
    if not s:
        return None
    if re.fullmatch(r"\d+", s):
        return int(s)
    return None


def entry_is_academic(entry: dict[str, Any]) -> bool:
    return entry_academic_num(entry) is not None


def entry_slot_title(entry: dict[str, Any]) -> str:
    """Краткое имя интервала для строк состояния (им. падеж после «Идёт» / «Закончился»)."""
    n = entry_academic_num(entry)
    if n is not None:
        return f"урок {n}"
    s = str(entry.get("lesson") or "").strip()
    return s if s else "Интервал"


def last_academic_lesson_ended_before(entries: list[dict[str, Any]], now_minutes: int) -> int | None:
    """Номер последнего академического урока, который уже закончился к моменту now_minutes."""
    last: int | None = None
    for e in entries:
        ln = entry_academic_num(e)
        if ln is None:
            continue
        if time_to_minutes(e["end"]) < now_minutes:
            last = ln
    return last


def next_academic_lesson_after_index(entries: list[dict[str, Any]], after_index: int) -> int | None:
    """Следующий номер академического урока среди записей с индексом > after_index."""
    for i in range(after_index + 1, len(entries)):
        ln = entry_academic_num(entries[i])
        if ln is not None:
            return ln
    return None


def display_name_for_academic_lesson_number(entries: list[dict[str, Any]], lesson_num: int | None) -> str:
    if lesson_num is None:
        return "следующего занятия"
    for e in entries:
        if entry_academic_num(e) == lesson_num:
            return entry_slot_title(e)
    return f"урок {lesson_num}"


def get_screen_bell_entries_and_bells(
    screen: dict[str, Any], target_date: date
) -> tuple[list[dict[str, Any]], dict[str, Any], str, dict[str, Any] | None]:
    """entries, bells, имя источника, активный шаблон или разовое переопределение (для last_lesson и т.п.)."""
    bells = load_bell_schedules()
    iso_date = target_date.isoformat()
    weekday = str(target_date.weekday())
    template_id = screen.get("bell_schedule_template") or "standard"

    date_override = next((item for item in bells.get("date_overrides", []) if item.get("date") == iso_date), None)
    if date_override:
        entries = date_override.get("entries", [])
        source_name = date_override.get("name", "Разовое")
        return entries, bells, source_name, date_override
    weekday_override_id = screen.get("weekday_bell_templates", {}).get(weekday) or bells.get("weekday_overrides", {}).get(weekday)
    if weekday_override_id:
        template_id = weekday_override_id
    template = next((item for item in bells.get("templates", []) if item.get("id") == template_id), None)
    if not template:
        template = bells.get("templates", [default_bell_schedules()["templates"][0]])[0]
    entries = template.get("entries", [])
    source_name = template.get("name", "Обычное")
    return entries, bells, source_name, template


def last_lesson_cap(template: dict[str, Any] | None) -> int | None:
    """Последний номер урока, после которого звонки «начало/конец» для больших номеров отключены; None — без ограничения."""
    if not template:
        return None
    raw = template.get("last_lesson")
    if raw is None:
        return None
    if isinstance(raw, str) and not str(raw).strip():
        return None
    try:
        n = int(raw)
    except (TypeError, ValueError):
        return None
    return max(1, min(24, n))


def resolve_bell_sound_url(filename: str | None) -> str | None:
    if not filename or not str(filename).strip():
        return None
    fn = str(filename).strip()
    if fn.startswith("/"):
        return fn
    p = BELL_SOUNDS_DIR / fn
    if p.is_file():
        return f"/uploads/bells/{fn}"
    p2 = UPLOADS_DIR / fn
    if p2.is_file():
        return f"/uploads/{fn}"
    return None


def bell_audio_url_to_path(url: str | None) -> Path | None:
    """Локальный путь к файлу звонка по URL из build_bell_audio_payload (только /uploads/…)."""
    if not url:
        return None
    u = str(url).strip()
    rest: str
    base: Path
    if u.startswith("/uploads/bells/"):
        rest = u[len("/uploads/bells/") :]
        base = BELL_SOUNDS_DIR
    elif u.startswith("/uploads/"):
        rest = u[len("/uploads/") :]
        base = UPLOADS_DIR
    else:
        return None
    name = Path(rest).name
    if not name:
        return None
    p = base / name
    return p if p.is_file() else None


def pick_entry_sound(ent: dict[str, Any], key: str, default_url: str | None) -> str | None:
    if key not in ent:
        return default_url
    raw = ent.get(key)
    if raw is None:
        return default_url
    s = str(raw).strip()
    if s in ("", "-", "none", "—"):
        return None
    return resolve_bell_sound_url(s) or default_url


def build_bell_audio_payload(
    screen: dict[str, Any],
    target_date: date,
    *,
    trigger_sec_window: int | None = None,
) -> dict[str, Any]:
    entries, bells, _, active_template = get_screen_bell_entries_and_bells(screen, target_date)
    cap_sched = compute_max_lesson_index_for_screen_day(screen, target_date)
    cap_tpl = last_lesson_cap(active_template)
    cap = cap_sched if cap_sched is not None else cap_tpl
    defaults = bells.get("sound_defaults") or {}
    d_start = resolve_bell_sound_url(defaults.get("start"))
    d_end = resolve_bell_sound_url(defaults.get("end"))
    rows = []
    for i, ent in enumerate(entries):
        n = entry_academic_num(ent)
        skip = cap is not None and n is not None and n > cap
        rows.append(
            {
                "index": i,
                "start": ent.get("start"),
                "end": ent.get("end"),
                "sound_start": None
                if skip
                else pick_entry_sound(ent, "sound_start", d_start),
                "sound_end": None if skip else pick_entry_sound(ent, "sound_end", d_end),
            }
        )
    tw = trigger_sec_window
    if tw is None:
        tw = DEFAULT_BELL_TRIGGER_SEC_WINDOW
    try:
        tw = int(tw)
    except (TypeError, ValueError):
        tw = DEFAULT_BELL_TRIGGER_SEC_WINDOW
    tw = max(5, min(55, tw))
    return {
        "defaults": {"start": d_start, "end": d_end},
        "entries": rows,
        "day": target_date.isoformat(),
        "trigger_sec_window": tw,
        "max_lesson_index_cap": cap,
        "max_lesson_index_from_schedule": cap_sched,
    }


def build_bell_status(screen: dict[str, Any], target_date: date) -> dict[str, Any]:
    entries, _bells, source_name, _tpl = get_screen_bell_entries_and_bells(screen, target_date)

    now = datetime.now()
    now_minutes = now.hour * 60 + now.minute
    status = {
        "template_name": source_name,
        "entries": entries,
        "current_lesson": None,
        "next_lesson": None,
        "state": "done",
        "message": "Занятия завершены",
        "countdown_text": "Звонков нет",
        "schedule_title": "Уроки закончились",
    }

    if not entries:
        status["message"] = "Расписание звонков не задано"
        return status

    first_start = time_to_minutes(entries[0]["start"])
    if now_minutes < first_start:
        diff = first_start - now_minutes
        status["state"] = "before"
        first = entries[0]
        cap = entry_slot_title(first)
        status["next_lesson"] = entry_academic_num(first)
        mp = minutes_phrase(diff)
        status["message"] = f"До {cap}, {mp}"
        status["countdown_text"] = f"Осталось {mp}"
        status["schedule_title"] = f"До {cap}, {mp}"
        return status

    for index, entry in enumerate(entries):
        start = time_to_minutes(entry["start"])
        end = time_to_minutes(entry["end"])
        if start <= now_minutes <= end:
            diff = end - now_minutes
            if not entry_is_academic(entry):
                last_ac = last_academic_lesson_ended_before(entries, now_minutes)
                next_ac = next_academic_lesson_after_index(entries, index)
                label = entry_slot_title(entry)
                mp = minutes_phrase(diff)
                status["state"] = "activity"
                status["current_lesson"] = last_ac
                status["next_lesson"] = next_ac
                status["message"] = f"Идёт {label}, до конца {mp}"
                status["countdown_text"] = f"До конца — {mp}"
                status["schedule_title"] = f"Идёт {label}, до конца {mp}"
                return status
            ln = entry_academic_num(entry)
            if ln is None:
                ln = 0
            shown = entry_slot_title(entry)
            mp = minutes_phrase(diff)
            status["state"] = "lesson"
            status["current_lesson"] = ln
            line = f"Идёт {shown}, до перемены {mp}"
            status["message"] = line
            status["countdown_text"] = f"До перемены {mp}"
            status["schedule_title"] = line
            return status
        next_start = time_to_minutes(entries[index + 1]["start"]) if index + 1 < len(entries) else None
        if end < now_minutes and next_start is not None and now_minutes < next_start:
            diff = next_start - now_minutes
            last_ac = last_academic_lesson_ended_before(entries, now_minutes)
            next_ac = next_academic_lesson_after_index(entries, index)
            ended_entry = entries[index]
            ended_label = entry_slot_title(ended_entry)
            mp = minutes_phrase(diff)
            status["state"] = "break"
            status["current_lesson"] = last_ac
            status["next_lesson"] = next_ac
            line = f"Закончился {ended_label}. Перемена {mp}"
            status["message"] = line
            status["countdown_text"] = f"Перемена {mp}"
            status["schedule_title"] = line
            return status

    status["countdown_text"] = "Занятия завершены"
    status["schedule_title"] = "Уроки закончились"
    return status


def tomorrow_schedule_visible(bell_status: dict[str, Any]) -> bool:
    """Блок «следующий учебный день» на ТВ — только после окончания последнего интервала дня (звонки завершены)."""
    return bell_status.get("state") == "done"


_WEEKDAY_ALIASES: dict[str, int] = {
    "понедельник": 0,
    "пн": 0,
    "mon": 0,
    "monday": 0,
    "вторник": 1,
    "вт": 1,
    "tue": 1,
    "tues": 1,
    "tuesday": 1,
    "среда": 2,
    "ср": 2,
    "wed": 2,
    "wednesday": 2,
    "четверг": 3,
    "чт": 3,
    "thu": 3,
    "thursday": 3,
    "пятница": 4,
    "пт": 4,
    "fri": 4,
    "friday": 4,
    "суббота": 5,
    "сб": 5,
    "sat": 5,
    "saturday": 5,
    "воскресенье": 6,
    "вс": 6,
    "sun": 6,
    "sunday": 6,
}


def parse_weekday_value(raw: Any) -> int | None:
    """Понедельник=0 … воскресенье=6, как :meth:`datetime.date.weekday`."""
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return int(raw.date().weekday())
    if isinstance(raw, date):
        return int(raw.weekday())
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        n = int(raw)
        if 0 <= n <= 6:
            return n
        if 1 <= n <= 7:
            return (n - 1) % 7
        return None
    s = str(raw).strip()
    if not s:
        return None
    if re.fullmatch(r"\d{1,2}", s):
        n = int(s)
        if 0 <= n <= 6:
            return n
        if 1 <= n <= 7:
            return (n - 1) % 7
        return None
    return _WEEKDAY_ALIASES.get(s.lower().replace("ё", "е"))


def parse_lesson_column_index(header: str) -> int | None:
    """Номер урока из заголовка столбца: «Урок1», «урок 7», «УРОК6» (без чувствительности к регистру)."""
    s = str(header or "").strip()
    if not s:
        return None
    sl = s.lower().replace("ё", "е")
    if not sl.startswith("урок"):
        return None
    tail = sl[4:].strip()
    if not tail:
        return None
    m = re.match(r"^(\d+)", tail)
    if not m:
        return None
    try:
        n = int(m.group(1))
    except ValueError:
        return None
    return n if 1 <= n <= 24 else None


def _pad_row_to_width(row: tuple[Any, ...] | list[Any] | None, width: int) -> list[Any]:
    if row is None:
        return [None] * width
    lst = list(row)
    if len(lst) < width:
        lst.extend([None] * (width - len(lst)))
    return lst[:width]


def _column_has_data_in_rows(data_rows: list[list[Any]], col_idx: int) -> bool:
    for row in data_rows:
        if col_idx >= len(row):
            continue
        v = row[col_idx]
        if v is None:
            continue
        if str(v).strip() != "":
            return True
    return False


def _discover_lesson_specs(
    headers: list[str],
    class_col_idx: int,
    data_rows: list[list[Any]],
) -> list[tuple[int, int]]:
    n = len(headers)
    explicit: list[tuple[int, int]] = []
    for j in range(class_col_idx + 1, n):
        lix = parse_lesson_column_index(headers[j])
        if lix is not None:
            explicit.append((j, lix))
    explicit.sort(key=lambda x: x[0])
    used_j = {j for j, _ in explicit}
    last_j = max((j for j, _ in explicit), default=class_col_idx)
    next_num = max((lix for _, lix in explicit), default=0) + 1

    inferred: list[tuple[int, int]] = []
    for j in range(last_j + 1, n):
        if j in used_j:
            continue
        if headers[j].strip() != "":
            break
        if _column_has_data_in_rows(data_rows, j):
            inferred.append((j, next_num))
            next_num += 1

    specs = explicit + inferred
    specs.sort(key=lambda x: x[0])
    return specs


def parse_weekly_schedule_excel(file_path: Path, *, lang: str = "ru") -> list[dict[str, Any]]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    if sheet.max_row is None or sheet.max_row < 2:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Файл Excel пуст.", "The Excel file is empty."),
        )
    max_col = int(sheet.max_column or 1)
    last_row = int(sheet.max_row)
    rows_raw = list(
        sheet.iter_rows(
            min_row=1,
            max_row=last_row,
            min_col=1,
            max_col=max_col,
            values_only=True,
        )
    )
    rows = [_pad_row_to_width(r, max_col) for r in rows_raw]
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]

    aliases = {"день недели", "день_недели", "weekday", "day of week"}
    weekday_col_idx: int | None = None
    for i, h in enumerate(headers):
        hl = str(h).strip().lower().replace("ё", "е")
        if hl in aliases:
            weekday_col_idx = i
            break
    if weekday_col_idx is None or "Класс" not in headers:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                lang,
                "Ожидаются колонки «День недели» (или Weekday) и «Класс».",
                "Expected columns «День недели» (or Weekday) and «Класс».",
            ),
        )

    class_col_idx = headers.index("Класс")
    data_rows = rows[1:]
    lesson_specs = _discover_lesson_specs(headers, class_col_idx, data_rows)
    if not lesson_specs:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                lang,
                "Нужны колонки 'Урок1', 'Урок2' и т.д.",
                "Lesson columns are required (e.g. Урок1, Урок2, …).",
            ),
        )

    result: list[dict[str, Any]] = []
    for row in data_rows:
        raw_wd = row[weekday_col_idx] if weekday_col_idx < len(row) else None
        raw_class = row[class_col_idx] if class_col_idx < len(row) else None
        wd = parse_weekday_value(raw_wd)
        if wd is None or not raw_class:
            continue

        lessons = []
        for j, lix in lesson_specs:
            value = row[j] if j < len(row) else None
            lessons.append({"index": lix, "subject": "" if value is None else str(value).strip()})

        result.append(
            {
                "weekday": wd,
                "class_name": str(raw_class).strip(),
                "class_key": normalize_class(str(raw_class)),
                "lessons": lessons,
            }
        )
    return result


def parse_excel(file_path: Path, *, lang: str = "ru") -> list[dict[str, Any]]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    if sheet.max_row is None or sheet.max_row < 2:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Файл Excel пуст.", "The Excel file is empty."),
        )
    max_col = int(sheet.max_column or 1)
    last_row = int(sheet.max_row)
    rows_raw = list(
        sheet.iter_rows(
            min_row=1,
            max_row=last_row,
            min_col=1,
            max_col=max_col,
            values_only=True,
        )
    )
    rows = [_pad_row_to_width(r, max_col) for r in rows_raw]
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    if "Дата" not in headers or "Класс" not in headers:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                lang,
                "Ожидаются колонки 'Дата' и 'Класс'.",
                "Expected columns 'Дата' and 'Класс'.",
            ),
        )

    class_col_idx = headers.index("Класс")
    date_col_idx = headers.index("Дата")
    data_rows = rows[1:]
    lesson_specs = _discover_lesson_specs(headers, class_col_idx, data_rows)
    if not lesson_specs:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                lang,
                "Нужны колонки 'Урок1', 'Урок2' и т.д.",
                "Lesson columns are required (e.g. Урок1, Урок2, …).",
            ),
        )

    result: list[dict[str, Any]] = []
    for row in data_rows:
        raw_date = row[date_col_idx] if date_col_idx < len(row) else None
        raw_class = row[class_col_idx] if class_col_idx < len(row) else None
        if not raw_date or not raw_class:
            continue

        if isinstance(raw_date, datetime):
            parsed_date = raw_date.date()
        elif isinstance(raw_date, date):
            parsed_date = raw_date
        else:
            parsed_date = datetime.strptime(str(raw_date).strip(), "%d.%m.%Y").date()

        lessons = []
        for j, lix in lesson_specs:
            value = row[j] if j < len(row) else None
            lessons.append({"index": lix, "subject": "" if value is None else str(value).strip()})

        result.append(
            {
                "date": parsed_date.isoformat(),
                "class_name": str(raw_class).strip(),
                "class_key": normalize_class(str(raw_class)),
                "lessons": lessons,
            }
        )
    return result


def parse_holidays_excel(file_path: Path, *, lang: str = "ru") -> list[dict[str, Any]]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Файл праздников Excel пуст.", "The holidays Excel file is empty."),
        )

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    required = {"Дата", "Название", "Описание"}
    if not required.issubset(set(headers)):
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                lang,
                "Ожидаются колонки 'Дата', 'Название', 'Описание'.",
                "Expected columns 'Дата', 'Название', 'Описание'.",
            ),
        )

    result: list[dict[str, Any]] = []
    for row in rows[1:]:
        if row is None:
            continue
        item = dict(zip(headers, row))
        raw_date = item.get("Дата")
        raw_name = item.get("Название")
        raw_description = item.get("Описание")
        if not raw_date or not raw_name:
            continue

        kind: str
        parsed_date: date | None = None
        md: str | None = None
        if isinstance(raw_date, datetime):
            parsed_date = raw_date.date()
            kind = "once"
        elif isinstance(raw_date, date):
            parsed_date = raw_date
            kind = "once"
        else:
            s0 = str(raw_date).strip()
            # допускаем пробелы вокруг точек и хвостовую точку: "01.01", "01. 01", "01.01.", "01 . 01 . 2026"
            s = s0.replace("\u00a0", " ")
            s = re.sub(r"\s+", "", s)
            s = s.rstrip(".")
            try:
                # dd.mm.yyyy -> once, dd.mm -> annual
                if re.fullmatch(r"\d{1,2}\.\d{1,2}\.\d{4}", s):
                    parsed_date = datetime.strptime(s, "%d.%m.%Y").date()
                    kind = "once"
                elif re.fullmatch(r"\d{1,2}\.\d{1,2}", s):
                    d_s, m_s = s.split(".")
                    md = f"{int(m_s):02d}-{int(d_s):02d}"
                    kind = "annual"
                else:
                    # неизвестный формат — пропускаем строку вместо падения ТВ
                    continue
            except (ValueError, TypeError):
                continue

        row_out: dict[str, Any] = {
            "name": str(raw_name).strip(),
            "description": "" if raw_description is None else str(raw_description).strip(),
            "kind": kind,
        }
        if kind == "annual" and md:
            row_out["md"] = md  # MM-DD
        elif parsed_date is not None:
            row_out["date"] = parsed_date.isoformat()
        else:
            continue
        result.append(row_out)
    # стабильная сортировка для хранения: annual по md, once по date
    def _k(item: dict[str, Any]) -> tuple[int, str]:
        if item.get("kind") == "annual":
            return (0, str(item.get("md") or ""))
        return (1, str(item.get("date") or "9999-12-31"))

    result.sort(key=_k)
    return result


def maybe_import_schedule_from_folder() -> None:
    signature = import_signature(AUTO_SCHEDULE_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("schedule")
    if current == signature and SCHEDULE_PATH.exists():
        return
    parsed = parse_excel(AUTO_SCHEDULE_IMPORT_PATH)
    write_json(SCHEDULE_PATH, parsed)
    state["schedule"] = signature
    save_import_state(state)


def maybe_import_full_schedule_from_folder() -> None:
    signature = import_signature(AUTO_FULL_SCHEDULE_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("full_schedule")
    if current == signature and FULL_SCHEDULE_PATH.exists():
        return
    parsed = parse_weekly_schedule_excel(AUTO_FULL_SCHEDULE_IMPORT_PATH)
    write_json(FULL_SCHEDULE_PATH, parsed)
    state["full_schedule"] = signature
    save_import_state(state)


def maybe_import_schedule_sample_from_folder() -> None:
    signature = import_signature(AUTO_SCHEDULE_SAMPLE_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("schedule_sample")
    if current == signature and SCHEDULE_SAMPLE_PATH.exists():
        return
    parsed = parse_weekly_schedule_excel(AUTO_SCHEDULE_SAMPLE_IMPORT_PATH)
    write_json(SCHEDULE_SAMPLE_PATH, parsed)
    state["schedule_sample"] = signature
    save_import_state(state)


def maybe_import_holidays_from_folder() -> None:
    signature = import_signature(AUTO_HOLIDAYS_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("holidays")
    if current == signature and HOLIDAYS_PATH.exists():
        return
    parsed = parse_holidays_excel(AUTO_HOLIDAYS_IMPORT_PATH)
    write_json(HOLIDAYS_PATH, parsed)
    state["holidays"] = signature
    save_import_state(state)


def parse_announcements_excel(file_path: Path, *, lang: str = "ru") -> list[dict[str, Any]]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Файл объявлений Excel пуст.", "The announcements Excel file is empty."),
        )

    # Берём первую колонку; заголовок может быть любым, пропускаем первую строку если похожа на заголовок.
    lines: list[str] = []
    for i, row in enumerate(rows):
        if row is None:
            continue
        v = row[0] if len(row) else None
        if v is None:
            s = ""
        else:
            s = str(v).strip()
        if i == 0 and s.lower() in {"текст", "объявление", "объявления", "announcement", "announcements"}:
            continue
        lines.append(s)

    blocks: list[list[str]] = [[]]
    for raw in lines:
        s = str(raw or "").rstrip()
        if not s:
            # пустые строки сохраняем как пустую строку внутри блока (для форматирования),
            # но не создаём бесконечные пустые блоки
            if blocks and blocks[-1]:
                blocks[-1].append("")
            continue
        st = s.strip()
        if st in ("!!!", "—!!!—"):
            if blocks and blocks[-1]:
                blocks.append([])
            continue
        if st.startswith("//") or st.startswith("#"):
            continue
        blocks[-1].append(s)

    out: list[dict[str, Any]] = []
    for b in blocks:
        text = "\n".join(b).strip("\n")
        if not text.strip():
            continue
        out.append({"text": text})
    return out


def maybe_import_announcements_from_folder() -> None:
    signature = import_signature(AUTO_ANNOUNCEMENTS_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("announcements")
    if current == signature and ANNOUNCEMENTS_PATH.exists():
        return
    parsed = parse_announcements_excel(AUTO_ANNOUNCEMENTS_IMPORT_PATH)
    write_json(ANNOUNCEMENTS_PATH, parsed)
    state["announcements"] = signature
    save_import_state(state)


def parse_marquee_excel(file_path: Path, *, lang: str = "ru") -> list[str]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Файл бегущей строки Excel пуст.", "The marquee Excel file is empty."),
        )
    out: list[str] = []
    for i, row in enumerate(rows):
        if row is None:
            continue
        v = row[0] if len(row) else None
        s = "" if v is None else str(v).strip()
        if i == 0 and s.lower() in {"текст", "бегущая строка", "marquee"}:
            continue
        if not s:
            continue
        if s.startswith("//") or s.startswith("#"):
            continue
        out.append(s)
    return out


def maybe_import_marquee_from_folder() -> None:
    signature = import_signature(AUTO_MARQUEE_IMPORT_PATH)
    if not signature:
        return
    state = load_import_state()
    current = state.get("marquee")
    if current == signature and MARQUEE_PATH.exists():
        return
    parsed = parse_marquee_excel(AUTO_MARQUEE_IMPORT_PATH)
    write_json(MARQUEE_PATH, parsed)
    state["marquee"] = signature
    save_import_state(state)


def _norm_subject_cell(value: str) -> str:
    return str(value or "").strip().lower()


def _subject_at_lesson(row: dict[str, Any] | None, lesson_index: int) -> str:
    if not row:
        return ""
    for les in row.get("lessons") or []:
        try:
            if int(les.get("index")) != lesson_index:
                continue
        except (TypeError, ValueError):
            continue
        sub = les.get("subject")
        return "" if sub is None else str(sub).strip()
    return ""


def _row_by_weekday_class(rows: list[dict[str, Any]], weekday: int, class_key: str) -> dict[str, Any] | None:
    hit: dict[str, Any] | None = None
    for item in rows:
        try:
            wd = int(item.get("weekday"))
        except (TypeError, ValueError):
            continue
        if wd == weekday and item.get("class_key") == class_key:
            hit = item
    return hit


def lesson_indices_union(
    dated_row: dict[str, Any] | None,
    full_row: dict[str, Any] | None,
    sample_row: dict[str, Any] | None,
) -> list[int]:
    idxs: set[int] = set()
    for src in (dated_row, full_row, sample_row):
        if not src:
            continue
        for les in src.get("lessons") or []:
            try:
                idxs.add(int(les.get("index")))
            except (TypeError, ValueError):
                continue
    return sorted(idxs) if idxs else list(range(1, 9))


def collect_enriched_schedule_rows(
    *,
    day: date,
    marker_reference_date: date,
    schedule_data: list[dict[str, Any]],
    full_data: list[dict[str, Any]],
    sample_data: list[dict[str, Any]],
    overrides: list[dict[str, Any]],
    selectors: list[str],
    bell_status: dict[str, Any],
) -> list[dict[str, Any]]:
    """Те же строки, что таблица уроков на ТВ: объединение дат/недели/образца и override по дате."""
    wd = day.weekday()
    day_iso = day.isoformat()
    dated_by_key = {
        item["class_key"]: item
        for item in schedule_data
        if item["date"] == day_iso and class_in_selected(item["class_key"], selectors)
    }
    keys_seen: set[str] = set(dated_by_key.keys())
    for item in full_data:
        try:
            if int(item.get("weekday")) != wd:
                continue
        except (TypeError, ValueError):
            continue
        if class_in_selected(item["class_key"], selectors):
            keys_seen.add(item["class_key"])
    for item in sample_data:
        try:
            if int(item.get("weekday")) != wd:
                continue
        except (TypeError, ValueError):
            continue
        if class_in_selected(item["class_key"], selectors):
            keys_seen.add(item["class_key"])

    display_names: dict[str, str] = {}
    for ck in keys_seen:
        dr = dated_by_key.get(ck)
        fr = _row_by_weekday_class(full_data, wd, ck)
        sr = _row_by_weekday_class(sample_data, wd, ck)
        display_names[ck] = str(
            (dr.get("class_name") if dr else None)
            or (fr.get("class_name") if fr else None)
            or (sr.get("class_name") if sr else None)
            or ck
        )

    sorted_keys = sorted(keys_seen, key=lambda k: class_sort_key(display_names[k]))
    apply_time_markers = day == marker_reference_date
    enriched: list[dict[str, Any]] = []

    for class_key in sorted_keys:
        dated_row = dated_by_key.get(class_key)
        full_row = _row_by_weekday_class(full_data, wd, class_key)
        sample_row = _row_by_weekday_class(sample_data, wd, class_key)
        indices = lesson_indices_union(dated_row, full_row, sample_row)
        lessons_out: list[dict[str, Any]] = []
        for lesson_index in indices:
            subject_full = _subject_at_lesson(full_row, lesson_index)
            subject_dated = _subject_at_lesson(dated_row, lesson_index) if dated_row else ""
            if dated_row is not None:
                base_subject = subject_dated if subject_dated != "" else subject_full
            else:
                base_subject = subject_full

            subject_sample = _subject_at_lesson(sample_row, lesson_index)
            use_sample = (
                sample_row is not None
                and subject_sample != ""
                and _norm_subject_cell(subject_sample) != _norm_subject_cell(subject_full)
            )
            if use_sample:
                cell_subject = subject_sample
                is_sample_diff = True
            else:
                cell_subject = base_subject
                is_sample_diff = False

            override = next(
                (
                    o
                    for o in overrides
                    if o["date"] == day_iso
                    and o["class_key"] == class_key
                    and int(o["lesson_index"]) == lesson_index
                ),
                None,
            )
            if override:
                final_subject = str(override["subject"]).strip()
                is_override = True
                is_sample_diff = False
            else:
                final_subject = cell_subject
                is_override = False

            cur = as_lesson_number(bell_status.get("current_lesson"))
            idx = as_lesson_number(lesson_index)
            lessons_out.append(
                {
                    "index": lesson_index,
                    "subject": final_subject,
                    "is_override": is_override,
                    "is_sample_diff": is_sample_diff,
                    "is_past": apply_time_markers and cur is not None and idx is not None and idx < cur,
                    "is_current": apply_time_markers and cur is not None and idx is not None and idx == cur,
                }
            )
        enriched.append({"class_name": display_names[class_key], "lessons": lessons_out})
    return enriched


def max_lesson_index_from_enriched_rows(today_rows: list[dict[str, Any]]) -> int | None:
    """Максимальный номер урока, для которого в сетке есть непустой предмет (по всем выбранным классам)."""
    best: int | None = None
    for row in today_rows:
        for les in row.get("lessons") or []:
            subj = str(les.get("subject") or "").strip()
            if not subj:
                continue
            try:
                ix = int(les.get("index"))
            except (TypeError, ValueError):
                continue
            best = ix if best is None else max(best, ix)
    return best


_ML_INDEX_CACHE: dict[tuple[Any, str], tuple[int | None, float]] = {}
_ML_INDEX_TTL_SEC = 12.0


def compute_max_lesson_index_for_screen_day(screen: dict[str, Any], day: date) -> int | None:
    """Последний урок дня по расписанию экрана (как в списке уроков); кэш на несколько секунд для ПК-цикла."""
    sid = screen.get("id")
    dk = day.isoformat()
    now = time.time()
    hit = _ML_INDEX_CACHE.get((sid, dk))
    if hit and now - hit[1] < _ML_INDEX_TTL_SEC:
        return hit[0]

    classes = screen.get("selected_classes", [])
    selectors = [normalize_class(item) for item in classes if normalize_class(item)]
    if not selectors:
        _ML_INDEX_CACHE[(sid, dk)] = (None, now)
        return None

    schedule_data = load_schedule()
    full_data = load_full_schedule()
    sample_data = load_schedule_sample()
    overrides = load_overrides()
    rows = collect_enriched_schedule_rows(
        day=day,
        marker_reference_date=date.min,
        schedule_data=schedule_data,
        full_data=full_data,
        sample_data=sample_data,
        overrides=overrides,
        selectors=selectors,
        bell_status={"current_lesson": None},
    )
    m = max_lesson_index_from_enriched_rows(rows)
    _ML_INDEX_CACHE[(sid, dk)] = (m, now)
    return m


def build_schedule_payload(screen: dict[str, Any], target_date: date) -> dict[str, Any]:
    schedule_data = load_schedule()
    full_data = load_full_schedule()
    sample_data = load_schedule_sample()
    overrides = load_overrides()
    bell_status = build_bell_status(screen, target_date)
    classes = screen.get("selected_classes", [])
    selectors = [normalize_class(item) for item in classes if normalize_class(item)]

    next_school_date_iso = None
    future_dates = sorted(
        {
            item["date"]
            for item in schedule_data
            if class_in_selected(item["class_key"], selectors) and item["date"] > target_date.isoformat()
        }
    )
    if future_dates:
        next_school_date_iso = future_dates[0]
    next_school_date = (
        date.fromisoformat(next_school_date_iso)
        if next_school_date_iso
        else date.fromordinal(target_date.toordinal() + 1)
    )
    show_next_day = tomorrow_schedule_visible(bell_status)

    today_rows = collect_enriched_schedule_rows(
        day=target_date,
        marker_reference_date=target_date,
        schedule_data=schedule_data,
        full_data=full_data,
        sample_data=sample_data,
        overrides=overrides,
        selectors=selectors,
        bell_status=bell_status,
    )
    tomorrow_rows = collect_enriched_schedule_rows(
        day=next_school_date,
        marker_reference_date=target_date,
        schedule_data=schedule_data,
        full_data=full_data,
        sample_data=sample_data,
        overrides=overrides,
        selectors=selectors,
        bell_status=bell_status,
    )
    max_lesson_index_today = max_lesson_index_from_enriched_rows(today_rows)

    return {
        "today": target_date.isoformat(),
        "next_school_day": next_school_date.isoformat(),
        "today_rows": today_rows,
        "tomorrow_rows": tomorrow_rows,
        "max_lesson_index_today": max_lesson_index_today,
        "tomorrow_schedule_visible": show_next_day,
        "bell_status": bell_status,
    }


ensure_dirs()

import bell_rupor_worker


@asynccontextmanager
async def _guard_school_lifespan(_app: FastAPI):
    bell_rupor_worker.start_worker()
    yield
    bell_rupor_worker.stop_worker()


app = FastAPI(title="GuardSchool", lifespan=_guard_school_lifespan)
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/uploads", StaticFiles(directory=UPLOADS_DIR), name="uploads")


@app.get("/api/version")
def get_public_version() -> dict[str, str]:
    return {"product": "GuardSchool", "version": APP_VERSION}


@app.get("/", response_class=HTMLResponse)
def root(request: Request) -> Response:
    if not load_auth():
        return RedirectResponse("/setup")
    if not is_authenticated(request):
        return RedirectResponse("/login")
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/setup", response_class=HTMLResponse)
def setup_page() -> Response:
    return FileResponse(STATIC_DIR / "setup.html")


@app.get("/login", response_class=HTMLResponse)
def login_page() -> Response:
    return FileResponse(STATIC_DIR / "login.html")


@app.get("/brand-logo")
def brand_logo() -> Response:
    logo_path = resolve_brand_logo_path()
    if not logo_path:
        raise HTTPException(status_code=404, detail="Логотип не найден (положите ico.png рядом с программой или в сборку).")
    return FileResponse(logo_path)


@app.get("/screen/{slug}", response_class=HTMLResponse)
def screen_page(slug: str) -> Response:
    # ТВ часто агрессивно кэширует HTML; без revalidate ссылка на styles.css может устареть при обновлении сборки.
    return FileResponse(
        STATIC_DIR / "screen.html",
        headers={"Cache-Control": "no-cache, must-revalidate"},
    )


@app.get("/api/bootstrap")
def bootstrap_state() -> dict[str, Any]:
    return {"configured": bool(load_auth())}


@app.post("/api/setup")
async def setup_admin(request: Request, username: str = Form(...), password: str = Form(...)) -> dict[str, str]:
    loc = admin_ui_lang(request)
    if load_auth():
        raise HTTPException(
            status_code=409,
            detail=admin_msg(loc, "Администратор уже создан.", "Administrator account already exists."),
        )
    if not password_is_valid(password):
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                loc,
                "Пароль: не менее 8 символов, нужны буквы и цифры.",
                "Password: at least 8 characters, with letters and digits.",
            ),
        )
    salt = secrets.token_hex(16)
    write_json(AUTH_PATH, {"username": username.strip(), "salt": salt, "password_hash": hash_password(password, salt)})
    return {"status": "ok"}


@app.post("/api/login")
async def login(request: Request, response: Response, username: str = Form(...), password: str = Form(...)) -> dict[str, str]:
    loc = admin_ui_lang(request)
    auth = load_auth()
    if not auth:
        raise HTTPException(
            status_code=428,
            detail=admin_msg(loc, "Сначала создайте администратора.", "Create the administrator account first."),
        )
    if username.strip() != auth["username"] or hash_password(password, auth["salt"]) != auth["password_hash"]:
        raise HTTPException(
            status_code=401,
            detail=admin_msg(loc, "Неверный логин или пароль.", "Invalid username or password."),
        )
    token = create_session_token(username.strip(), auth)
    sec = session_cookie_secure(request)
    response.set_cookie(
        SESSION_COOKIE,
        token,
        httponly=True,
        samesite="lax",
        secure=sec,
        path="/",
    )
    return {"status": "ok"}


@app.post("/api/logout")
def logout(request: Request, response: Response) -> dict[str, str]:
    sec = session_cookie_secure(request)
    response.delete_cookie(SESSION_COOKIE, path="/", secure=sec, samesite="lax")
    return {"status": "ok"}


@app.get("/api/admin/config")
def get_admin_config(request: Request) -> dict[str, Any]:
    require_auth(request)
    return load_config()


@app.post("/api/admin/config")
async def save_admin_config(request: Request) -> dict[str, str]:
    require_auth(request)
    payload = await request.json()
    write_json(CONFIG_PATH, sanitize_config(payload))
    return {"status": "ok"}


@app.get("/api/admin/background-gallery")
def admin_background_gallery(request: Request, folder: str = "") -> dict[str, Any]:
    require_auth(request)
    f = _safe_rel_uploads_subdir(folder)
    return {
        "folder": f,
        "folders": list_background_subdirs_from_uploads(),
        "images": list_background_images_from_uploads(f),
    }


@app.post("/api/admin/screen-bg-next")
async def admin_screen_bg_next(request: Request) -> dict[str, Any]:
    require_auth(request)
    body = await request.json()
    sid = str(body.get("screen_id") or "").strip()
    if not sid:
        raise HTTPException(status_code=400, detail="Нужен screen_id.")
    config = load_config()
    screen = next((s for s in config.get("screens") or [] if str(s.get("id")) == sid), None)
    if not screen:
        raise HTTPException(status_code=404, detail="Экран не найден.")
    # сбрасываем «фиксированный» фон, если был
    screen["background_force_image"] = ""
    screen["background_rotate_enabled"] = True
    try:
        cur = int(screen.get("background_rotate_cursor", 0))
    except (TypeError, ValueError):
        cur = 0
    screen["background_rotate_cursor"] = max(0, cur) + 1
    screen["background_rotate_epoch"] = int(time.time())
    write_json(CONFIG_PATH, sanitize_config(config))
    return {"status": "ok", "screen_id": sid}


async def _handle_pc_audio_test_play(request: Request) -> dict[str, str]:
    """Тест колонок с ПК: первый файл из uploads/bells или явный filename в JSON."""
    require_auth(request)
    body: dict[str, Any] = {}
    raw = await request.body()
    if raw.strip():
        try:
            parsed = json.loads(raw.decode("utf-8-sig"))
            if isinstance(parsed, dict):
                body = parsed
        except (json.JSONDecodeError, UnicodeError, TypeError):
            body = {}
    use_first = request.query_params.get("use_first") in ("1", "true", "yes")
    raw_fn = body.get("filename")
    if raw_fn in (None, "", False):
        raw_fn = body.get("file")
    if use_first or raw_fn in (None, "", False):
        fn = ""
    else:
        fn = str(raw_fn).strip()
    if fn:
        safe = Path(fn).name
        path = BELL_SOUNDS_DIR / safe
        if not path.is_file():
            path = UPLOADS_DIR / safe
        if not path.is_file():
            raise HTTPException(status_code=404, detail="Файл не найден в каталогах uploads/bells или uploads.")
    else:
        first = first_bell_sound_path()
        if not first:
            raise HTTPException(
                status_code=400,
                detail="В uploads/bells нет файлов. Загрузите звук в разделе «Звонки».",
            )
        path = first
        safe = path.name
    audio = dict(sanitize_audio_stream(load_config().get("audio_stream")))
    vol = body.get("volume_percent")
    if vol is not None:
        try:
            audio["volume_percent"] = max(0, min(100, int(vol)))
        except (TypeError, ValueError):
            pass
    ok, err = bell_rupor_worker.stream_file_to_rupor(path, audio)
    if not ok:
        raise HTTPException(status_code=400, detail=err or "Не удалось запустить отправку.")
    return {
        "status": "ok",
        "detail": f"Файл «{safe}»: воспроизведение на ПК в фоне (если ffmpeg с muxer wasapi — через него; иначе ffplay в той же папке). Ниже — stderr.",
    }


@app.post("/api/admin/audio-stream-test-send")
async def post_audio_stream_test_send(request: Request) -> dict[str, str]:
    """Совместимость со старыми клиентами; предпочтительно /api/admin/pc-audio-test-play."""
    return await _handle_pc_audio_test_play(request)


@app.post("/api/admin/pc-audio-test-play")
async def post_pc_audio_test_play(request: Request) -> dict[str, str]:
    """Тест Рупор из сайдбара (есть только в актуальном сервере — иначе в браузере будет 404)."""
    return await _handle_pc_audio_test_play(request)


@app.get("/api/admin/audio-stream-last-send")
def get_audio_stream_last_send(request: Request) -> dict[str, Any]:
    """Последний сеанс ffplay (команда и stderr)."""
    require_auth(request)
    return {"last": bell_rupor_worker.get_last_ffmpeg_send()}


@app.get("/api/admin/audio-stream-status")
def get_audio_stream_status(request: Request) -> dict[str, Any]:
    """Процесс ffplay и последний результат — для панели админки."""
    require_auth(request)
    return bell_rupor_worker.get_ffmpeg_status()


@app.post("/api/admin/audio-stream-stop")
def post_audio_stream_stop(request: Request) -> dict[str, Any]:
    """Остановить текущее воспроизведение (ffplay)."""
    require_auth(request)
    return bell_rupor_worker.stop_current_stream()


@app.get("/api/admin/break-music-files")
def list_break_music_files(request: Request) -> dict[str, Any]:
    """Файлы в data/break_music + порядок воспроизведения и громкости для админки."""
    require_auth(request)
    ensure_dirs()
    files = []
    if BREAK_MUSIC_DIR.exists():
        for p in sorted(BREAK_MUSIC_DIR.iterdir()):
            if p.is_file():
                files.append({"filename": p.name})
    import local_audio_worker

    return {
        "files": files,
        "directory": str(BREAK_MUSIC_DIR),
        "playback": local_audio_worker.get_break_music_playback_info(),
    }


@app.post("/api/admin/break-music-preview")
async def post_break_music_preview(request: Request) -> dict[str, str]:
    """Прослушать файл из data/break_music на ПК с громкостью «перемена» (как в оркестрации)."""
    require_auth(request)
    body: dict[str, Any] = {}
    raw = await request.body()
    if raw.strip():
        try:
            parsed = json.loads(raw.decode("utf-8-sig"))
            if isinstance(parsed, dict):
                body = parsed
        except (json.JSONDecodeError, UnicodeError, TypeError):
            body = {}
    fn = Path(str(body.get("filename") or "")).name
    if not fn:
        raise HTTPException(status_code=400, detail="Укажите filename (имя файла в data/break_music).")
    path = BREAK_MUSIC_DIR / fn
    if not path.is_file():
        raise HTTPException(status_code=404, detail="Файл не найден в data/break_music.")
    import local_audio_worker

    audio = sanitize_audio_stream(load_config().get("audio_stream"))
    vol = body.get("volume_percent")
    if vol is not None:
        try:
            audio = {**audio, "break_music_volume_percent": max(0, min(100, int(vol)))}
        except (TypeError, ValueError):
            pass
    local_audio_worker.stop_playback_hard()
    time.sleep(0.08)
    ok, err = local_audio_worker.play_break_music_preview(path, audio)
    if not ok:
        raise HTTPException(status_code=400, detail=err or "Не удалось воспроизвести.")
    v = int(audio.get("break_music_volume_percent") or 40)
    return {
        "status": "ok",
        "detail": f"«{fn}» — громкость перемены {v}% (ffmpeg af volume={v/100.0:.2f}).",
    }


@app.post("/api/admin/upload-background")
async def upload_background(request: Request, file: UploadFile = File(...)) -> dict[str, str]:
    require_auth(request)
    suffix = Path(file.filename or "background").suffix or ".jpg"
    target = UPLOADS_DIR / f"{secrets.token_hex(8)}{suffix}"
    target.write_bytes(await file.read())
    return {"path": f"/uploads/{target.name}"}


@app.post("/api/admin/upload-widget-image")
async def upload_widget_image(request: Request, file: UploadFile = File(...)) -> dict[str, str]:
    require_auth(request)
    ensure_dirs()
    sub = UPLOADS_DIR / WIDGET_IMAGES_SUBDIR
    sub.mkdir(parents=True, exist_ok=True)
    suffix = Path(file.filename or "image").suffix or ".jpg"
    target = sub / f"{secrets.token_hex(8)}{suffix}"
    target.write_bytes(await file.read())
    return {"path": f"/uploads/{WIDGET_IMAGES_SUBDIR}/{target.name}"}


@app.post("/api/admin/upload-holidays")
async def upload_holidays(request: Request, file: UploadFile = File(...)) -> dict[str, Any]:
    require_auth(request)
    lang = admin_ui_lang(request)
    suffix = Path(file.filename or "holidays.xlsx").suffix or ".xlsx"
    temp = UPLOADS_DIR / f"holidays_import{suffix}"
    temp.write_bytes(await file.read())
    parsed = parse_holidays_excel(temp, lang=lang)
    write_json(HOLIDAYS_PATH, parsed)
    return {"status": "ok", "rows": len(parsed)}


@app.post("/api/admin/upload-announcements")
async def upload_announcements(request: Request, file: UploadFile = File(...)) -> dict[str, Any]:
    require_auth(request)
    lang = admin_ui_lang(request)
    suffix = Path(file.filename or "announcements.xlsx").suffix or ".xlsx"
    temp = UPLOADS_DIR / f"announcements_import{suffix}"
    temp.write_bytes(await file.read())
    parsed = parse_announcements_excel(temp, lang=lang)
    write_json(ANNOUNCEMENTS_PATH, parsed)
    return {"status": "ok", "rows": len(parsed)}


@app.post("/api/admin/upload-marquee")
async def upload_marquee(request: Request, file: UploadFile = File(...)) -> dict[str, Any]:
    require_auth(request)
    lang = admin_ui_lang(request)
    suffix = Path(file.filename or "marquee.xlsx").suffix or ".xlsx"
    temp = UPLOADS_DIR / f"marquee_import{suffix}"
    temp.write_bytes(await file.read())
    parsed = parse_marquee_excel(temp, lang=lang)
    write_json(MARQUEE_PATH, parsed)
    return {"status": "ok", "rows": len(parsed)}


@app.post("/api/admin/upload-schedule")
async def upload_schedule(request: Request, file: UploadFile = File(...)) -> dict[str, Any]:
    require_auth(request)
    lang = admin_ui_lang(request)
    suffix = Path(file.filename or "schedule.xlsx").suffix or ".xlsx"
    temp = UPLOADS_DIR / f"schedule_import{suffix}"
    temp.write_bytes(await file.read())
    parsed = parse_excel(temp, lang=lang)
    write_json(SCHEDULE_PATH, parsed)
    return {"status": "ok", "rows": len(parsed)}


@app.post("/api/admin/upload-full-schedule")
async def upload_full_schedule(request: Request, file: UploadFile = File(...)) -> dict[str, Any]:
    require_auth(request)
    lang = admin_ui_lang(request)
    suffix = Path(file.filename or "full_schedule.xlsx").suffix or ".xlsx"
    temp = UPLOADS_DIR / f"full_schedule_import{suffix}"
    temp.write_bytes(await file.read())
    parsed = parse_weekly_schedule_excel(temp, lang=lang)
    write_json(FULL_SCHEDULE_PATH, parsed)
    return {"status": "ok", "rows": len(parsed)}


@app.post("/api/admin/upload-schedule-sample")
async def upload_schedule_sample(request: Request, file: UploadFile = File(...)) -> dict[str, Any]:
    require_auth(request)
    lang = admin_ui_lang(request)
    suffix = Path(file.filename or "schedule_sample.xlsx").suffix or ".xlsx"
    temp = UPLOADS_DIR / f"schedule_sample_import{suffix}"
    temp.write_bytes(await file.read())
    parsed = parse_weekly_schedule_excel(temp, lang=lang)
    write_json(SCHEDULE_SAMPLE_PATH, parsed)
    return {"status": "ok", "rows": len(parsed)}


@app.post("/api/admin/upload-bell-sound")
async def upload_bell_sound(request: Request, file: UploadFile = File(...)) -> dict[str, str]:
    require_auth(request)
    lang = admin_ui_lang(request)
    suffix = (Path(file.filename or "sound").suffix or ".mp3").lower()
    if suffix not in {".mp3", ".wav", ".ogg", ".m4a", ".aac"}:
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Допустимы: mp3, wav, ogg, m4a, aac.", "Allowed: mp3, wav, ogg, m4a, aac."),
        )
    name = f"{secrets.token_hex(6)}{suffix}"
    target = BELL_SOUNDS_DIR / name
    target.write_bytes(await file.read())
    return {"filename": name, "url": f"/uploads/bells/{name}"}


@app.get("/api/admin/bell-sounds")
def list_bell_sounds(request: Request) -> dict[str, Any]:
    require_auth(request)
    ensure_dirs()
    files = []
    if BELL_SOUNDS_DIR.exists():
        for p in sorted(BELL_SOUNDS_DIR.iterdir()):
            if p.is_file():
                files.append({"filename": p.name, "url": f"/uploads/bells/{p.name}"})
    return {"files": files}


@app.get("/api/admin/schedule")
def get_schedule_snapshot(request: Request) -> dict[str, Any]:
    require_auth(request)
    return {
        "schedule": load_schedule(),
        "holidays": load_holidays(),
        "announcements": load_announcements(),
        "marquee": load_marquee_items(),
        "overrides": load_overrides(),
        "bells": load_bell_schedules(),
        "history": load_change_log(),
        "full_schedule_rows": len(load_full_schedule()),
        "schedule_sample_rows": len(load_schedule_sample()),
        "imports": {
            "schedule_path": str(AUTO_SCHEDULE_IMPORT_PATH),
            "full_schedule_path": str(AUTO_FULL_SCHEDULE_IMPORT_PATH),
            "schedule_sample_path": str(AUTO_SCHEDULE_SAMPLE_IMPORT_PATH),
            "holidays_path": str(AUTO_HOLIDAYS_IMPORT_PATH),
            "announcements_path": str(AUTO_ANNOUNCEMENTS_IMPORT_PATH),
            "marquee_path": str(AUTO_MARQUEE_IMPORT_PATH),
            "state": load_import_state(),
        },
        "app_version": APP_VERSION,
    }


@app.get("/api/admin/history")
def get_change_history(request: Request) -> dict[str, Any]:
    require_auth(request)
    return {"history": load_change_log(), "app_version": APP_VERSION}


@app.post("/api/admin/overrides")
async def save_overrides(request: Request) -> dict[str, str]:
    require_auth(request)
    payload = await request.json()
    write_json(OVERRIDES_PATH, payload)
    return {"status": "ok"}


@app.post("/api/admin/bells")
async def save_bells(request: Request) -> dict[str, str]:
    require_auth(request)
    payload = await request.json()
    write_json(BELL_SCHEDULES_PATH, payload)
    return {"status": "ok"}


@app.get("/api/admin/export")
def export_admin_bundle(request: Request) -> StreamingResponse:
    require_auth(request)
    bundle = export_bundle_bytes()
    filename = f'gorniitv_export_{datetime.now().strftime("%Y%m%d_%H%M%S")}.zip'
    return StreamingResponse(
        io.BytesIO(bundle),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/admin/export-weekly-schedule")
def export_weekly_schedule_bundle(request: Request) -> StreamingResponse:
    require_auth(request)
    bundle = export_weekly_schedule_bundle_bytes()
    filename = f'guardschool_weekly_schedule_{datetime.now().strftime("%Y%m%d_%H%M%S")}.zip'
    return StreamingResponse(
        io.BytesIO(bundle),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/admin/weekly-schedule-template.xlsx")
def download_weekly_schedule_template_xlsx(request: Request) -> FileResponse:
    require_auth(request)
    ensure_dirs()
    ensure_weekly_schedule_template_file()
    if not FULL_SCHEDULE_SAMPLE_XLSX.is_file():
        raise HTTPException(
            status_code=404,
            detail=admin_msg(admin_ui_lang(request), "Не удалось создать шаблон.", "Could not create template."),
        )
    return FileResponse(
        FULL_SCHEDULE_SAMPLE_XLSX,
        filename="full_schedule_sample.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.post("/api/admin/import")
async def import_admin_bundle(request: Request, file: UploadFile = File(...)) -> dict[str, str]:
    require_auth(request)
    lang = admin_ui_lang(request)
    if not (file.filename or "").lower().endswith(".zip"):
        raise HTTPException(
            status_code=400,
            detail=admin_msg(lang, "Нужен ZIP-архив экспорта.", "Expected an export ZIP archive."),
        )
    import_bundle_bytes(await file.read(), lang=lang)
    return {"status": "ok"}


@app.post("/api/admin/import-weekly-schedule")
async def import_weekly_schedule_bundle(request: Request, file: UploadFile = File(...)) -> dict[str, Any]:
    require_auth(request)
    lang = admin_ui_lang(request)
    if not (file.filename or "").lower().endswith(".zip"):
        raise HTTPException(
            status_code=400,
            detail=admin_msg(
                lang,
                "Нужен ZIP с full_schedule.json и/или schedule_sample.json.",
                "Expected a ZIP with full_schedule.json and/or schedule_sample.json.",
            ),
        )
    import_weekly_schedule_bundle_bytes(await file.read(), lang=lang)
    return {
        "status": "ok",
        "full_schedule_rows": len(load_full_schedule()),
        "schedule_sample_rows": len(load_schedule_sample()),
    }


@app.post("/api/admin/preview-payload")
def post_preview_payload(request: Request, payload: dict[str, Any] = Body(...)) -> JSONResponse:
    require_auth(request)
    screen = payload.get("screen")
    if not isinstance(screen, dict):
        raise HTTPException(status_code=400, detail="Нужен объект screen.")
    screen.setdefault("selected_classes", [])
    screen.setdefault("bell_schedule_template", "standard")
    screen.setdefault("weekday_bell_templates", {})
    screen.setdefault("widgets", [])
    cfg = load_config()
    today = calendar_today_for_config(cfg)
    audio = sanitize_audio_stream(cfg.get("audio_stream"))
    import local_audio_worker

    body = {
        "schedule": build_schedule_payload(screen, today),
        "holidays": load_holidays(),
        "announcements": load_announcements(),
        "marquee": load_marquee_items(),
        "background_gallery": list_background_images_from_uploads(screen.get("background_rotate_folder")),
        "bell_audio": build_bell_audio_payload(
            screen,
            today,
            trigger_sec_window=audio_trigger_sec_window(audio),
        ),
        "pc_audio_preview": local_audio_worker.describe_pc_audio_preview(
            screen,
            cfg.get("audio_stream"),
        ),
        "display": display_settings_dict(cfg),
    }
    return JSONResponse(content=body)


@app.get("/api/screen/{slug}")
def get_screen(slug: str) -> JSONResponse:
    config = load_config()
    screen = next((item for item in config["screens"] if item["slug"] == slug and item.get("is_active", True)), None)
    if not screen:
        raise HTTPException(status_code=404, detail="Экран не найден.")
    today = calendar_today_for_config(config)
    audio = sanitize_audio_stream(config.get("audio_stream"))
    payload = {
        "screen": screen,
        "serverTime": datetime.now().isoformat(),
        "schedule": build_schedule_payload(screen, today),
        "holidays": load_holidays(),
        "announcements": load_announcements(),
        "marquee": load_marquee_items(),
        "background_gallery": list_background_images_from_uploads(screen.get("background_rotate_folder")),
        "bell_audio": build_bell_audio_payload(
            screen,
            today,
            trigger_sec_window=audio_trigger_sec_window(audio),
        ),
        "display": display_settings_dict(config),
    }
    return JSONResponse(
        content=payload,
        headers={
            "Cache-Control": "no-store, no-cache, must-revalidate, max-age=0",
            "Pragma": "no-cache",
            "Expires": "0",
        },
    )
